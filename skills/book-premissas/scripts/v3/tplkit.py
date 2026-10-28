# Kit para gerar books a partir do template v3 (Campeche). Copiado do repo da Rachel Souto
# (rachelsouto-source/book-premissas-template/scripts/tplkit.py); unica mudanca: finish() aceita o texto do rodape.
import copy, io, os, hashlib
from pptx import Presentation
from pptx.util import Inches, Emu
from pptx.oxml.ns import qn
from PIL import Image

E = 914400
STATUS_COLOR = {"DEFINIDO": "1E7F5C", "A VALIDAR": "FC6058", "PENDENTE": "C98200", "A COTAR": "5B6B8C"}
STATUS_ICON = {"DEFINIDO": "✓", "A VALIDAR": "!", "PENDENTE": "?", "A COTAR": "$"}
KEEP_TAGS = {qn("p:sp"), qn("p:pic"), qn("p:grpSp"), qn("p:cxnSp"), qn("p:graphicFrame")}


class Book:
    def __init__(self, tpl_path, cache_dir):
        self.prs = Presentation(tpl_path)
        self.tpl = list(self.prs.slides)
        self.n_tpl = len(self.tpl)
        self.cache = cache_dir
        os.makedirs(cache_dir, exist_ok=True)
        self._img_open = {}
        # pills de status (slide 24 do template)
        s24 = self.tpl[23]
        self.pill_xml = {}
        for sh in s24.shapes:
            if sh.has_text_frame and sh.text_frame.text in STATUS_COLOR and sh.shape_type == 1:
                self.pill_xml.setdefault(sh.text_frame.text, sh._element)

    # ---------- slides ----------
    def dup(self, tpl_no):
        src = self.tpl[tpl_no - 1]
        s = self.prs.slides.add_slide(src.slide_layout)
        for shp in list(s.shapes):
            shp._element.getparent().remove(shp._element)
        tree = s.shapes._spTree
        for el in src.shapes._spTree:
            if el.tag in KEEP_TAGS:
                tree.append(copy.deepcopy(el))
        # copiar fundo
        bg = src._element.cSld.find(qn("p:bg"))
        if bg is not None:
            s._element.cSld.insert(0, copy.deepcopy(bg))
        # remapear imagens
        for blip in tree.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            part = src.part.related_part(rid)
            blip.set(qn("r:embed"), s.part.relate_to(part, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"))
        return Slide(self, s)

    def finish(self, out, footer_name, footer_mid="BOOK DE PREMISSAS DE FACHADA E INTERIORES DAS ÁREAS COMUNS"):
        # remover slides originais do template
        sldIdLst = self.prs.slides._sldIdLst
        for sldId in list(sldIdLst)[: self.n_tpl]:
            self.prs.part.drop_rel(sldId.get(qn("r:id")))
            sldIdLst.remove(sldId)
        for n, sl in enumerate(self.prs.slides, 1):
            for cnv in sl.shapes._spTree.iter(qn("p:cNvPr")):
                d = cnv.get("descr", "")
                if "/" in d or "\\" in d:
                    cnv.set("descr", os.path.basename(d.replace("\\", "/")))
            for sh in sl.shapes:
                if sh.has_text_frame and "BOOK DE PREMISSAS DE FACHADA E INTERIORES  ·  " in sh.text_frame.text and "SPOT" in sh.text_frame.text:
                    set_paras(sh, [f"{footer_name}  ·  {footer_mid}  ·  {n:02d}"])
        self.prs.save(out)

    # ---------- imagens ----------
    def img(self, path):
        if path not in self._img_open:
            im = Image.open(path)
            im.load()
            self._img_open[path] = im.convert("RGB")
        return self._img_open[path]

    def window(self, path, aspect, focus=(0.5, 0.5), hfrac=None, zoom=1.0):
        """janela normalizada (x0,y0,x1,y1) com o aspecto do quadro"""
        im = self.img(path)
        W, H = im.size
        if hfrac is not None:
            wh = hfrac * H
        else:
            wh = min(H, W / aspect) * zoom
        ww = wh * aspect
        if ww > W:
            ww = W
            wh = ww / aspect
        cx, cy = focus[0] * W, focus[1] * H
        x0 = min(max(cx - ww / 2, 0), W - ww)
        y0 = min(max(cy - wh / 2, 0), H - wh)
        return (x0 / W, y0 / H, (x0 + ww) / W, (y0 + wh) / H)

    def crop_file(self, path, win, frame_w_in, dpi=220):
        im = self.img(path)
        W, H = im.size
        box = (round(win[0] * W), round(win[1] * H), round(win[2] * W), round(win[3] * H))
        key = hashlib.md5(f"{path}{box}{frame_w_in:.2f}".encode()).hexdigest()[:12]
        out = os.path.join(self.cache, key + ".jpg")
        if not os.path.exists(out):
            c = im.crop(box)
            tw = int(frame_w_in * dpi)
            if c.size[0] > tw:
                c = c.resize((tw, round(c.size[1] * tw / c.size[0])), Image.LANCZOS)
            c.save(out, quality=86)
        return out


def set_paras(sh, texts):
    """troca o texto mantendo o estilo de cada parágrafo (1 run por parágrafo no template)"""
    txBody = sh.text_frame._txBody
    ps = txBody.findall(qn("a:p"))
    while len(ps) < len(texts):
        new = copy.deepcopy(ps[-1])
        ps[-1].addnext(new)
        ps = txBody.findall(qn("a:p"))
    for p in ps[len(texts):]:
        txBody.remove(p)
    ps = txBody.findall(qn("a:p"))
    for p, t in zip(ps, texts):
        runs = p.findall(qn("a:r"))
        if not runs:
            continue
        for r in runs[1:]:
            p.remove(r)
        runs[0].find(qn("a:t")).text = t


class Slide:
    def __init__(self, book, s):
        self.b = book
        self.s = s
        self._next_id = 5000

    def sh(self, sid):
        for x in self.s.shapes:
            if x.shape_id == sid:
                return x
        raise KeyError(sid)

    def T(self, sid, *texts):
        set_paras(self.sh(sid), list(texts))
        return self

    def color(self, sid, hexrgb):
        sp = self.sh(sid)._element.spPr
        for c in sp.findall(qn("a:solidFill")):
            c.find(qn("a:srgbClr")).set("val", hexrgb)

    def font_color(self, sid, hexrgb):
        for r in self.sh(sid)._element.iter(qn("a:rPr")):
            sf = r.find(qn("a:solidFill"))
            if sf is not None and sf.find(qn("a:srgbClr")) is not None:
                sf.find(qn("a:srgbClr")).set("val", hexrgb)

    def delete(self, *sids):
        for sid in sids:
            el = self.sh(sid)._element
            el.getparent().remove(el)

    def geom(self, sid, x=None, y=None, w=None, h=None):
        x_ = self.sh(sid)
        if x is not None: x_.left = Emu(int(x * E))
        if y is not None: x_.top = Emu(int(y * E))
        if w is not None: x_.width = Emu(int(w * E))
        if h is not None: x_.height = Emu(int(h * E))

    def pos(self, sid):
        x = self.sh(sid)
        return x.left / E, x.top / E, x.width / E, x.height / E

    def clone(self, sid, dx=0, dy=0, x=None, y=None):
        el = copy.deepcopy(self.sh(sid)._element)
        self._next_id += 1
        nid = self._next_id
        cnv = el.find(".//" + qn("p:cNvPr"))
        cnv.set("id", str(nid))
        cnv.set("name", f"Shape {nid}")
        self.sh(sid)._element.addnext(el)
        sh = self.sh(nid)
        if x is not None: sh.left = Emu(int(x * E))
        else: sh.left = Emu(int(sh.left + dx * E))
        if y is not None: sh.top = Emu(int(y * E))
        else: sh.top = Emu(int(sh.top + dy * E))
        return nid

    def pill(self, status, x, y):
        el = copy.deepcopy(self.b.pill_xml[status])
        self._next_id += 1
        nid = self._next_id
        cnv = el.find(".//" + qn("p:cNvPr"))
        cnv.set("id", str(nid)); cnv.set("name", f"Pill {nid}")
        self.s.shapes._spTree.append(el)
        sh = self.sh(nid)
        sh.left = Emu(int(x * E)); sh.top = Emu(int(y * E))
        return nid

    def status_icon(self, sid, status):
        self.color(sid, STATUS_COLOR[status])
        set_paras(self.sh(sid), [STATUS_ICON[status]])

    def P(self, sid, path, focus=(0.5, 0.5), hfrac=None, zoom=1.0):
        """troca a imagem preenchendo o quadro (recorte real); devolve a janela usada"""
        pic = self.sh(sid)
        aspect = pic.width / pic.height
        win = self.b.window(path, aspect, focus, hfrac, zoom)
        f = self.b.crop_file(path, win, pic.width / E)
        _, rid = self.s.part.get_or_add_image_part(f)
        blip = pic._element.blipFill.find(qn("a:blip"))
        old = blip.get(qn("r:embed"))
        blip.set(qn("r:embed"), rid)
        src = pic._element.blipFill.find(qn("a:srcRect"))
        if src is not None:
            for k in ("l", "r", "t", "b"):
                if k in src.attrib: del src.attrib[k]
        pic._element.find(".//" + qn("p:cNvPr")).set("descr", os.path.basename(path))
        # drop_rel do python-pptx só conta r:id, não r:embed: confere se outro
        # quadro do slide (ex.: um clone) ainda usa a imagem antiga antes de remover.
        if old != rid and not self.s._element.xpath(f'.//a:blip[@r:embed="{old}"]'):
            self.s.part.drop_rel(old)
        return win

    def map_pt(self, sid, win, u, v):
        x, y, w, h = self.pos(sid)
        fx = (u - win[0]) / (win[2] - win[0])
        fy = (v - win[1]) / (win[3] - win[1])
        return x + fx * w, y + fy * h, (0 <= fx <= 1 and 0 <= fy <= 1)
