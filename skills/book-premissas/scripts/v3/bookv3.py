"""Funcoes comuns para montar books no layout v3 (template Campeche) com o tplkit.
Base: build_canasvieiras.py da Rachel Souto. Ids de shapes = template v3 do repo dela.
"""
import os
from pptx import Presentation
from pptx.util import Pt
from tplkit import Book, E  # noqa: F401

REPO = os.path.expanduser("~/Claude/projetos/book-premissas-template")
TPL = os.path.join(REPO, "template", "TEMPLATE_v3_CAMPECHE_SPOT_Book_de_Premissas_rev13-10-2026.pptx")
D_ = "DEFINIDO"
LEG = "Status do item: DEFINIDO"

# ids do template (slide do template -> elementos)
V11 = [(321, 322, 323), (324, 325, 326), (327, 328, 329), (330, 331, 332), (333, 334, 335), (336, 337, 338),
       (339, 340, 341), (342, 343, 344), (345, 346, 347)]
V12 = [(371, 372, 373), (374, 375, 376), (377, 378, 379), (380, 381, 382), (383, 384, 385),
       (386, 387, 388), (389, 390, 391), (392, 393, 394), (395, 396, 397)]
V13 = [(430, 431, 432), (433, 434, 435), (436, 437, 438), (439, 440, 441), (442, 443, 444),
       (445, 446, 447), (448, 449, 450), (451, 452, 453), (454, 455, 456), (457, 458, 459)]
TEX = [(227, 228, 229), (231, 232, 233), (235, 236, 237), (239, 240, 241)]
MC = [(1067, 1068, 1069), (1071, 1072, 1073), (1075, 1076, 1077), (1079, 1080, 1081)]
MD = [(1148, 1149, 1150), (1152, 1153, 1154), (1156, 1157, 1158), (1160, 1161, 1162)]
CARDS15 = [(483, 484, 485, 486), (488, 489, 490, 491), (493, 494, 495, 496), (498, 499, 500, 501), (503, 504, 505, 506)]
ROWS27 = [(859, 860), (861, 862), (865, 866), (867, 868), (870, 871), (873, 874), (876, 877), (878, 879), (881, 882)]
PILLS27 = [863, 872]
ROWS30 = [(949, 950), (952, 953), (955, 956), (957, 958), (960, 961), (962, 963), (965, 966), (967, 968), (970, 971)]
PILLS30 = [951]
ROWS32 = [(1018, 1019), (1021, 1022), (1024, 1025), (1026, 1027), (1029, 1030), (1031, 1032), (1034, 1035), (1037, 1038), (1041, 1042)]
PILLS32 = [1020, 1036, 1039]


def old_markers(legacy_pptx, slide_no, pic_id=11):
    """posicoes normalizadas (u,v) dos marcadores numerados do book legado, na imagem inteira"""
    p = Presentation(legacy_pptx)
    sl = p.slides[slide_no - 1]
    pic = [s for s in sl.shapes if s.shape_id == pic_id][0]
    cl, cr, ct, cb = pic.crop_left, pic.crop_right, pic.crop_top, pic.crop_bottom
    out = {}
    for s in sl.shapes:
        if s.has_text_frame and s.text_frame.text.strip().isdigit() and s.width < 0.26 * E and \
           pic.left <= s.left <= pic.left + pic.width and pic.top <= s.top <= pic.top + pic.height:
            cx = s.left + s.width / 2; cy = s.top + s.height / 2
            u = cl + (cx - pic.left) / pic.width * (1 - cl - cr)
            v = ct + (cy - pic.top) / pic.height * (1 - ct - cb)
            out[int(s.text_frame.text)] = (u, v)
    return [out[k] for k in sorted(out)]


def place_markers(S, pic_id, win, marker_ids, pts):
    for mid, (u, v) in zip(marker_ids, pts):
        x, y, ok = S.map_pt(pic_id, win, u, v)
        w = S.pos(mid)[2]
        S.geom(mid, x - w / 2, y - w / 2)
        if not ok:
            print("  marcador fora do quadro:", mid)
    for mid in marker_ids[len(pts):]:
        S.delete(mid)


def label_w(S, sid, txt):
    S.T(sid, txt)
    S.geom(sid, w=0.30 + 0.078 * len(txt))


def card(S, ids, code, title, onde, spec, obra, img, focus=(0.5, 0.5), hfrac=None):
    pic, lab, txt = ids
    S.P(pic, img, focus, hfrac)
    S.T(lab, code)
    S.T(txt, title, "ONDE", onde, "ESPECIFICAÇÃO", spec, "ATENÇÃO NA OBRA", obra)


def rows(S, rws, data, pill_ids, full_w=5.98):
    for pid in pill_ids:
        S.delete(pid)
    for (lid, tid), (lab, txt) in zip(rws, data):
        S.T(lid, lab); S.T(tid, txt)
        S.geom(tid, w=full_w)
    for (lid, tid) in rws[len(data):]:
        S.delete(lid, tid)


def checklist(S, data):
    rws = [(775, 776, 777), (778, 779, 780), (782, 783, 784), (785, 786, 787), (789, 790, 791), (792, 793, 794)]
    for (box, tid, pid), (title, body) in zip(rws, data):
        S.T(tid, title, body)
        S.T(box, "✓")
        y = S.pos(pid)[1]
        S.delete(pid)
        S.pill(D_, 12.71 - 0.90, y)


def vista_items(S, pairs, data):
    for (nid, tid, iid), (title, body) in zip(pairs, data):
        S.T(tid, title, body)
        S.status_icon(iid, D_)
    for (nid, tid, iid) in pairs[len(data):]:
        S.delete(nid, tid, iid)


def concept_cards(S, data):
    """data = [(nome, img, focus, conceito, decisao)] ate 5"""
    for (pic, ttl, body, pill), (name, img, foc, conc, sit) in zip(CARDS15, data):
        S.P(pic, img, foc)
        S.T(ttl, name)
        S.T(body, "CONCEITO SEAZONE (SEGUIR)", conc, "DECISÃO REGISTRADA", sit)
        x, y, w, h = S.pos(pill)
        S.delete(pill)
        S.pill(D_, x, y)
    for ids in CARDS15[len(data):]:
        S.delete(*ids)


def finish(B, out, spot):
    """salva e reduz a fonte do rodape nas aberturas de secao (caixa estreita)"""
    B.finish(out, spot)
    P = Presentation(out)
    for sl in P.slides:
        for sh in sl.shapes:
            if sh.has_text_frame and "SPOT  ·  BOOK DE PREMISSAS" in sh.text_frame.text and sh.width < 5 * E:
                for para in sh.text_frame.paragraphs:
                    for r in para.runs:
                        r.font.size = Pt(5.5)
    P.save(out)
    print("ok", out)
