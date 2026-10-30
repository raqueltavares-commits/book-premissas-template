"""Cachoeira Spot no layout v3 (template Campeche). Conteudo = decisoes aprovadas na v2 (13/10/2026).
python build_cachoeira.py DADOS saida.pptx   (DADOS = trabalho/v3_dados/cachoeira)
"""
import os
import sys

sys.path.insert(0, os.path.expanduser("~/Claude/.claude/skills/book-premissas/scripts/v3"))
from bookv3 import *  # noqa: F401,F403
from pptx.util import Pt  # noqa: E402

D, OUT = sys.argv[1], sys.argv[2]
OLD = os.path.join(REPO, "books", "legado", "CACHOEIRA_SPOT_Book_de_Premissas_legado.pptx")
src = lambda n: os.path.join(D, "src", n + ".jpg")
sw = lambda n: os.path.join(D, "sw", n)
E = {k: src(k) for k in ["E01", "E02", "E03", "E04", "E05", "E06", "E07", "E08", "E09", "DRONE"]}
REC = lambda n: src(f"REC0{n}")
RF = lambda n: src(f"ROOF0{n}")
OLDI = {"brise": "s05_12_b242da24.jpg", "forro": "s05_19_80b104cd.jpg", "gc": "s05_26_6b3c10e1.jpg", "pilar": "s05_33_53cd9443.jpg",
        "logo": "s05_40_899485e6.jpg", "esq": "s05_47_71423ac8.jpg", "flor": "s06_12_ac1f6a4f.jpg", "letreiro": "s06_19_95a5527f.jpg",
        "base": "s06_26_0f436c8f.jpg", "perg": "s06_33_e42957e9.jpg", "grelha": "s06_40_72eb356a.jpg", "escada": "s06_47_128e1786.jpg",
        "moledo": "s19_12_43923a36.jpg", "cimento": "s19_19_f086bf56.jpg", "ripas": "s19_26_6486db51.jpg", "verde": "s19_33_82151144.jpg",
        "bancada": "s19_40_3b2f8e89.jpg", "piscina": "s19_47_e163c01b.jpg", "saogabriel": "s20_40_9801004e.jpg"}
O = lambda k: os.path.join(D, "old", OLDI[k])

B = Book(TPL, os.path.join(D, "cache"))

# 1 · CAPA
S = B.dup(1)
S.P(16, E["E05"], (0.5, 0.5))
S.P(17, E["E01"], (0.5, 0.5))
S.P(18, REC(1), (0.5, 0.5))
S.P(19, RF(7), (0.5, 0.5))
S.T(21, "BOOK DE PREMISSAS DE FACHADA E INTERIORES DAS ÁREAS COMUNS")
S.sh(21).text_frame.paragraphs[0].runs[0].font.size = Pt(8.5)
S.T(22, "Cachoeira Spot")
S.sh(22).text_frame.paragraphs[0].runs[0].font.size = Pt(36)
S.T(24, "Lançamento  ×  Executivo  ·  Fachada  ·  Interiores das áreas comuns  ·  Estilo Industrial")
S.T(25, "Rua da Amizade, 89 · Cachoeira do Bom Jesus · Florianópolis/SC", "DOCUMENTO TÉCNICO INTERNO")

# 2 · CONTEXTUALIZAÇÃO
S = B.dup(2)
S.T(36, "Consolidar as premissas de fachada e de interiores das áreas comuns do Cachoeira Spot a partir da leitura cruzada "
        "entre as imagens de lançamento, o pré-executivo da Bloco B e o projeto de interiores da Seazone: o que a imagem "
        "vendeu, o que o projeto especifica, onde o projeto precisa ser ajustado e o que a obra deve seguir na compra e na execução.")
S.P(37, E["E09"], (0.45, 0.5), 0.6)
S.P(40, O("brise"))
S.P(43, E["E08"], (0.35, 0.72), 0.25)
S.T(45, "CONFIRA", "o que a obra precisa observar e seguir")
for sid, q in zip([49, 51, 53, 55, 57, 59, 61], [
        "O que as imagens de lançamento mostram na fachada?",
        "Qual material o projeto define para cada elemento?",
        "Onde o projeto precisa ser ajustado à imagem vendida?",
        "Como são os brises, o forro amadeirado e os guarda-corpos?",
        "Qual é o conceito de interiores das áreas comuns que deve ser seguido?",
        "O que não pode ser alterado sem validação?",
        "O que a obra precisa observar na execução?"]):
    S.T(sid, q)
S.T(62, "Regra do material: premissas validadas pela Seazone. Estilo de interiores: Industrial. Prevalecem a imagem de lançamento, "
        "o conceito Seazone e o Briefing; o projeto compatibiliza. Fontes (1.20 - [3330] Cachoeira Spot / 02 - Projetos): imagens 3D, "
        "pré-executivo, Memorial Descritivo V02, projeto de interiores, Conferência de Projetos e Premissas de Projeto. "
        "Interiores das unidades, índices urbanísticos, áreas e custos não fazem parte deste Book.")

# 3 · COMO USAR
S = B.dup(3)
S.P(75, O("forro"))
S.T(77, "BV01")
S.T(78, "Código", "Código do projeto/memorial (ex.: BR01, BV01, J01, T07) para achar o item no projeto.")
S.P(87, E["E09"], (0.45, 0.5), None, 0.8)
S.P(90, REC(1), (0.5, 0.5))
S.P(93, RF(6), (0.5, 0.5))
S.T(89, "Fachada", "imagens de lançamento · paleta · leitura geral · texturas T01–T12 · vistas frontal, esquina e térreo · pontos de atenção")
S.T(92, "Interiores das áreas comuns", "conceito × decisão · recepção e corredores · rooftop interno e externo · pontos de atenção")
S.T(95, "Materiais de interiores das áreas comuns", "texturas T13–T24 para compra e aprovação de amostras, com MDFs e metais do estilo Industrial")
S.T(79, "STATUS DOS ITENS")
S.T(81, "premissa validada pela Seazone: seguir na compra e na execução")
S.delete(82, 83, 84, 85)

# 4 · DIVISOR 01
S = B.dup(4)
S.P(101, E["E06"], (0.62, 0.5))
S.T(105, "Imagens de lançamento  ·  paleta  ·  leitura geral  ·  texturas  ·  vistas com marcação e especificação  ·  pontos de atenção")

# 5 · MOODBOARD
S = B.dup(5)
S.P(118, E["E09"], (0.4, 0.5)); label_w(S, 119, "09 · ESQUINA")
S.P(120, E["E04"], (0.5, 0.5)); label_w(S, 121, "04 · FRONTAL")
S.P(122, E["E05"], (0.6, 0.5)); label_w(S, 123, "05")
S.P(124, E["DRONE"], (0.55, 0.5)); label_w(S, 125, "DRONE")
S.P(126, E["E08"], (0.35, 0.65)); label_w(S, 127, "08 · TÉRREO")
S.P(128, RF(7), (0.5, 0.5)); label_w(S, 129, "ROOFTOP")
for sid, (a, b) in zip([132, 134, 136, 138, 140], [
        ("Lajes amadeiradas", "bordas e forro das sacadas em ACM madeira, desenhando camadas horizontais"),
        ("Brises cor madeira", "ripas verticais em alumínio amadeirado em frente às janelas e porta-janelas"),
        ("Guarda-corpo de vidro", "sacadas em vidro com alumínio anodizado de montante aparente"),
        ("Base comercial", "lojas com vitrines, ripado escuro, marquise com LED e letreiro CACHOEIRA SPOT"),
        ("Coroamento", "floreira contínua, volume cinza com logo SEAZONE e pergolado preto no rooftop")]):
    S.T(sid, a, b)

# 6 · PALETA
S = B.dup(6)
for pic, lab, txt, code, a, b, img in [
        (151, 152, 153, "BV01–BV05", "Brise alumínio amadeirado", "frente das janelas", O("brise")),
        (154, 155, 156, "ACM MADEIRA", "Forro e bordas de laje", "sacadas e pavimentos", O("forro")),
        (157, 158, 159, "PILARES", "Papel de Presente B208", "pilares claros (como na imagem)", O("pilar")),
        (160, 161, 162, "GUARDA-CORPO", "Vidro + alumínio anodizado", "sacadas · montante aparente", O("gc")),
        (163, 164, 165, "ESQUADRIAS", "Linha 32 preto fosco", "janelas e porta-janelas", O("esq")),
        (166, 167, 168, "PERGOLADO", "Metal preto RAL 9005", "terraço do rooftop", O("perg"))]:
    S.P(pic, img); S.T(lab, code); S.T(txt, a, b)
S.T(170, "COMO LER", "Fachada = lajes com forro e bordas amadeirados (ACM madeira) + brises verticais em alumínio cor madeira + "
                     "guarda-corpo de vidro + pilares claros (Papel de Presente) e empenas cinza médio + esquadria preta. "
                     "Detalhe de cada material nas páginas de texturas T01–T12.")

# 7 · LEITURA GERAL
S = B.dup(7)
win = S.P(181, E["E09"])
label_w(S, 182, "LANÇAMENTO · 09 (ESQUINA)")
place_markers(S, 181, win, list(range(183, 193)), old_markers(OLD, 4))
S.T(193, "A fachada vendida é lida em camadas horizontais: lajes com bordas e forro amadeirados, sacadas com guarda-corpo de vidro "
         "e brises verticais cor madeira em frente às janelas. No topo, floreira contínua e volume cinza com o logo SEAZONE; na base, "
         "floreira sobre a marquise com o letreiro CACHOEIRA SPOT e lojas com ripado escuro.")
for sid, (a, b) in zip(range(196, 215, 2), [
        ("Volume superior com logo", "cinza, letreiro SEAZONE iluminado"),
        ("Floreira da cobertura", "vegetação pendente em toda a borda"),
        ("Guarda-corpo das sacadas", "vidro com alumínio anodizado, montante aparente"),
        ("Brise vertical amadeirado", "ripas em alumínio em frente às janelas"),
        ("Forro e bordas de laje", "acabamento amadeirado (ACM madeira)"),
        ("Pilares e paredes claras", "Papel de Presente Suvinil (B208) · empenas cinza médio"),
        ("Floreira do 1º pavimento", "faixa contínua sobre a marquise"),
        ("Letreiro CACHOEIRA SPOT", "letras brancas iluminadas sobre a marquise"),
        ("Base comercial", "ripado escuro entre vitrines das lojas"),
        ("Acesso", "escada com LED nos degraus")]):
    S.T(sid, a, b)

# 8-10 · TEXTURAS
NOTE = ("Amostras extraídas das imagens de lançamento: referência visual, não substitui amostra física. "
        "Especificação: pré-executivo (elevações 011 a 013), Memorial Descritivo V02 e Briefing.")
for title, cards in [
    ("Texturas da fachada · brises, forro e guarda-corpos (1/3)", [
        ("T01", "Brise das janelas", "Frente das janelas e porta-janelas de todas as fachadas.",
         "Brise em alumínio com acabamento amadeirado, ripas verticais cor madeira como na imagem.",
         "Executivo indica alumínio: manter, com acabamento amadeirado no tom da imagem; aprovar amostra antes da compra.", O("brise")),
        ("T02", "Forro e bordas amadeirados", "Face inferior das lajes das sacadas e bordas dos pavimentos.",
         "Forro de madeira em ACM madeira com perfis de vinco (prancha 017), inclinado na elevação 02.",
         "Converge com a imagem: aprovar o padrão da madeira do ACM com amostra antes da compra.", O("forro")),
        ("T03", "Guarda-corpo das sacadas", "Sacadas dos pavimentos e cobertura.",
         "Vidro com alumínio anodizado de montante aparente (Briefing) · piscina com perfil \"U\" inox.",
         "Detalhar perfil e fixação no executivo antes do pedido.", O("gc")),
        ("T04", "Pilares e paredes", "Pilares entre janelas, paredes e empenas.",
         "Pilares claros em Papel de Presente Suvinil (B208) e empenas em cinza médio, como na imagem · textura/massa acrílica.",
         "Projeto não define cor: seguir esta premissa e aprovar painel-teste antes da compra.", O("pilar"))]),
    ("Texturas da fachada · esquadrias, logo e vegetação (2/3)", [
        ("T05", "Volume superior + logo SEAZONE", "Caixa d'água / volume técnico da cobertura.",
         "Volume cinza com logo SEAZONE iluminado (frente) e casinha (fundos), como na imagem.",
         "Projetos detalha fixação e ponto elétrico antes do fechamento da fachada.", O("logo")),
        ("T06", "Esquadrias", "Janelas e porta-janelas das unidades.",
         "Linha 32 (gold) preto fosco (Briefing) · tabelas J01–J08 e P01–P16.",
         "Completar linha e cor nas tabelas de esquadrias antes do pedido.", O("esq")),
        ("T07", "Floreiras de fachada", "Borda da cobertura e faixa sobre a marquise do térreo.",
         "Floreiras A a I impermeabilizadas (membrana elastomérica) · Guaimbé e Russélia branca (paisagismo).",
         "Irrigação detalhada no executivo; teste de estanqueidade antes do plantio.", O("flor")),
        ("T08", "Brise das condensadoras", "Faixas nas lajes da fachada posterior/lateral.",
         "Brise para condensadoras em lâminas horizontais, cinza escuro como na imagem (elevação 02).",
         "Integrar com a conferência de climatização.", O("grelha"))]),
    ("Texturas da fachada · térreo e cobertura (3/3)", [
        ("T09", "Letreiro CACHOEIRA SPOT", "Frente da marquise do térreo.",
         "Letras brancas iluminadas, como na imagem (memorial: nome do edifício na fachada).",
         "Projetos detalha fixação e elétrica antes do fechamento da fachada.", O("letreiro")),
        ("T10", "Base comercial", "Fachada das lojas no térreo.",
         "Painéis ripados escuros entre as vitrines e marquise com linha de LED, como na imagem.",
         "Executar como na imagem; incluir ripado e LED nas elevações.", O("base")),
        ("T11", "Pergolado da cobertura", "Terraço do rooftop.",
         "Pergolado metálico preto RAL 9005, escuro como na imagem (estilo Industrial).",
         "Elevação 02 indica branco e o memorial corten/madeira: compatibilizar antes da fabricação.", O("perg")),
        ("T12", "Acesso", "Entrada principal pela calçada.",
         "Escada com LED nos degraus e guarda-corpo de vidro, como na imagem · acesso sem rampa no lobby.",
         "Escada, LED e guarda-corpo detalhados no executivo.", O("escada"))])]:
    S = B.dup(8)
    S.T(221, title); S.T(225, NOTE)
    for ids, c in zip(TEX, cards):
        card(S, ids, *c)

# 11 · VISTA FRONTAL (tpl 13)
S = B.dup(13)
S.T(410, "Fachada frontal (Rua da Amizade)")
win = S.P(414, E["E04"])
label_w(S, 415, "LANÇAMENTO · 04 (FRONTAL)")
place_markers(S, 414, win, list(range(416, 426)), old_markers(OLD, 7))
S.T(427, "LEITURA DA FACHADA FRONTAL", "A frente é uma grelha regular de sacadas com vidro e brises verticais amadeirados, emoldurada por "
         "lajes com forro de madeira. O brise segue a imagem vendida (ripas em alumínio cor madeira) e o projeto compatibiliza.")
S.T(429, LEG)
vista_items(S, V13, [
    ("Logo SEAZONE", "Iluminado no volume superior, como na imagem"),
    ("Floreira da cobertura", "Membrana elastomérica · Guaimbé e Russélia"),
    ("Brises", "Alumínio amadeirado, ripas como na imagem"),
    ("Guarda-corpo", "Vidro com alumínio anodizado de montante aparente (Briefing)"),
    ("Forro das lajes", "Forro de madeira em ACM madeira"),
    ("Pilares", "Papel de Presente Suvinil (B208), como na imagem"),
    ("Floreira do 1º pavimento", "Faixa contínua sobre a marquise"),
    ("Letreiro", "CACHOEIRA SPOT iluminado, como na imagem"),
    ("Lojas", "Vitrines + ripado escuro, como na imagem"),
    ("Acesso", "Escada com LED, como na imagem")])

# 12 · VISTA ESQUINA (tpl 12)
S = B.dup(12)
S.T(354, "Esquina e fachada lateral")
win = S.P(358, E["E05"])
label_w(S, 359, "LANÇAMENTO · 05")
place_markers(S, 358, win, list(range(360, 369)), old_markers(OLD, 8))
S.T(370, LEG)
vista_items(S, V12, [
    ("Floreira da cobertura", "Vegetação pendente na borda · membrana elastomérica"),
    ("Volume superior", "Cinza, com logo · reservatório com acabamento conforme fachada"),
    ("Bordas de laje", "Revestidas em madeira (ACM madeira), como na imagem"),
    ("Guarda-corpo", "Vidro com alumínio anodizado de montante aparente (Briefing)"),
    ("Brises", "Alumínio amadeirado, ripas como na imagem"),
    ("Paredes claras", "Papel de Presente Suvinil (B208) · textura acrílica"),
    ("Fachada lateral", "Sacadas repetidas com brise e vidro · brise das condensadoras"),
    ("Muro", "Muro claro (imagem) com logo e nome Seazone (Briefing)")])
S.P(398, O("forro")); S.P(399, O("brise")); S.P(400, O("pilar"))
S.T(401, "ACM · BRISE · PILAR")
S.T(403, "NAS ELEVAÇÕES", "As elevações 011 a 013 trazem só notas de material (brise, vidro, forro, pergolado). Sem paginação, cores "
         "nem caderno de fachada: vale a imagem de venda como referência de acabamento.")

# 13 · VISTA TÉRREO (tpl 13) — marcadores do legado (faixa 04 (2) = parte de baixo da imagem 04)
S = B.dup(13)
S.T(410, "Térreo: lojas, marquise e acesso")
win = S.P(414, E["E04"], (0.49, 0.85), 0.54)
label_w(S, 415, "LANÇAMENTO · 04 (TÉRREO)")
pts = [(u, 0.598 + v * 0.399) for u, v in old_markers(OLD, 9)]
place_markers(S, 414, win, list(range(416, 426)), pts)
S.T(427, "ATENÇÃO · IDENTIDADE DO ACESSO", "Letreiro, marquise com LED, ripado das lojas e escada iluminada são a \"cara\" do térreo "
         "no lançamento: executar como na imagem e incluir nas elevações.")
S.T(429, LEG)
vista_items(S, V13, [
    ("Letreiro", "CACHOEIRA SPOT em letras brancas iluminadas"),
    ("Floreira", "Faixa contínua de vegetação sobre a marquise"),
    ("Marquise", "Linha de LED na face inferior"),
    ("Ripado", "Painéis ripados escuros entre as vitrines"),
    ("Vitrines", "Esquadrias pretas, vidro incolor"),
    ("Guarda-corpo", "Vidro na frente das lojas"),
    ("Escada", "Degraus com LED embutido"),
    ("Canteiros", "Vegetação tropical baixa · Ipê-amarelo, Liriope (paisagismo)"),
    ("Mureta", "Mureta cinza com canteiro na frente das unidades"),
    ("Unidades garden", "Esquadrias com vegetação à frente")])

# 14 · PONTOS DE ATENÇÃO FACHADA
S = B.dup(24)
S.T(767, "01 · FACHADA · PONTOS DE ATENÇÃO")
S.T(768, "Fachada: pontos de atenção na obra")
S.P(772, E["E05"], (0.5, 0.5))
label_w(S, 773, "LANÇAMENTO · 05")
checklist(S, [
    ("Brises", "Alumínio amadeirado, ripas como na imagem: compatibilizar a cor das elevações (cinza) e aprovar amostra antes de fabricar."),
    ("Pergolado", "Metal preto RAL 9005, escuro como na imagem: compatibilizar elevação 02 (branco) e memorial (corten/madeira)."),
    ("Pilares e paredes", "Pilares Papel de Presente (B208) e empenas cinza médio, como na imagem: painel-teste antes da compra."),
    ("Esquadrias", "Linha 32 (gold) preto fosco (Briefing): completar as tabelas J01–J08 e P01–P16 antes do pedido."),
    ("Letreiros e logo", "Executar como na imagem; prever fixação e pontos elétricos antes do fechamento da fachada."),
    ("Base comercial e acesso", "Ripado escuro, marquise com LED e escada iluminada, como na imagem: incluir nas elevações.")])

# 15 · DIVISOR 02
S = B.dup(14)
S.P(465, REC(1), (0.5, 0.5))
S.T(468, "Interiores das áreas comuns")
S.T(469, "Conceito Seazone × decisão  ·  recepção e corredores  ·  rooftop interno  ·  rooftop externo e piscina  ·  pontos de atenção")

# 16 · CONCEITO × DECISÃO
S = B.dup(15)
S.T(477, "02 · INTERIORES DAS ÁREAS COMUNS")
S.T(478, "Conceito Seazone × decisão registrada")
concept_cards(S, [
    ("Recepção", REC(2), (0.5, 0.5),
     "Parede em pedra Moledo com logo CACHOEIRA SPOT, cimento queimado com frase, sofá orgânico off white, ripas de madeira",
     "Pedra Moledo + textura cimento queimado areia (Briefing)"),
    ("Corredores", REC(3), (0.5, 0.5),
     "Paredes brancas, portas com número preto e arandela em cada porta",
     "Pintura branca, arandelas, forro modular (Briefing)"),
    ("Rooftop interno", RF(3), (0.5, 0.5),
     "Copa com bancada em granito preto, balcão em madeira, pergolado de vidro",
     "Forro modular (Briefing) · balcão em MDF Carvalho Treviso, armários em MDF Chumbo"),
    ("Rooftop externo", RF(6), (0.5, 0.5),
     "Parede em pedra Moledo com logo, parede verde com balanço, banco de concreto, piscina com borda",
     "Reforço metálico do balanço · piscina esverdeada 20×20 detalhada no executivo"),
    ("Projeto de interiores", RF(1), (0.5, 0.5),
     "Imagens e memorial de interiores Seazone (10 - Projeto de Interiores)",
     "Sem executivo de interiores contratado: vale o conceito Seazone e este book")])
S.T(507, "Imagens e memorial de interiores Seazone × LOG de considerações, Briefing e Conferência de Projetos.")
S.T(509, "ATENÇÃO · REGRA PARA A OBRA", "Não há projeto executivo de interiores contratado. Vale o conceito Seazone (10 - Projeto de "
         "Interiores / 02 - Imagens e memorial descritivo de interiores).")

# 17 · FICHA RECEPÇÃO
S = B.dup(27)
S.T(848, "02 · INTERIORES DAS ÁREAS COMUNS · RECEPÇÃO E CORREDORES")
S.T(849, "Ficha: recepção, hall e corredores")
S.P(853, REC(1), (0.5, 0.5)); label_w(S, 854, "CONCEITO · RECEPÇÃO")
S.P(855, REC(3), (0.5, 0.5)); label_w(S, 856, "CONCEITO · CORREDOR")
S.T(857, "Frase da recepção: \"Seja você sua maior expectativa\", em letras pretas sobre o cimento queimado.")
rows(S, ROWS27, [
    ("Parede-destaque", "Pedra Moledo com logo CACHOEIRA SPOT iluminado (LOG e Briefing) · ponto elétrico para o logo"),
    ("Segunda parede", "Textura cimento queimado areia com a frase em letras pretas"),
    ("Piso", "Porcelanato 90×90 natural/acetinado · rodapé branco de poliestireno de sobrepor, quina viva"),
    ("Forro", "Gesso acartonado na recepção · forro modular nos corredores (Briefing)"),
    ("Paredes", "Pintura acrílica acabamento acetinado · corredores brancos"),
    ("Mobiliário", "Sofá orgânico off white, poltronas brancas, puffs cinza, mesa lateral preta (memorial de interiores)"),
    ("Decoração", "Nome do Spot preto, capacho Seazone azul, artes de marketing com moldura preta, números pretos"),
    ("Iluminação", "Recepção: 2 pendentes + spots de embutir pretos · corredores: arandela em cada porta, luz 4000 K"),
    ("Portas", "Brancas com bandeira até o forro · acesso em perfil metálico + vidro com reconhecimento facial")], PILLS27)

# 18 · MOODBOARD ROOFTOP (tpl 28)
S = B.dup(28)
S.T(888, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP · MOODBOARD")
S.T(889, "Rooftop: lazer coberto, copa e piscina")
for pic, img in zip(range(893, 902), [RF(1), RF(3), RF(2), RF(4), RF(5), RF(6), RF(7), src("ROOF_PLANTA"), E["E02"]]):
    S.P(pic, img)
S.T(902, "CONCEITO", "Balcão em madeira voltado para o mar com banquetas pretas, copa com bancada em granito São Gabriel e "
         "armários escuros, pergolado de vidro e piscina com borda.",
         "PAREDES", "Pedra Moledo com logo CACHOEIRA SPOT e a frase \"Seu lugar fora de casa\"; parede verde com balanço.",
         "ESTILO", "Industrial: MDF Carvalho Treviso e Chumbo, metal preto RAL 9005.")
for pic, lab, name, img in [(903, 904, "Moledo", O("moledo")), (905, 906, "Ripas", O("ripas")),
                            (907, 908, "Bancada", O("bancada")), (909, 910, "Piscina", O("piscina"))]:
    S.P(pic, img); S.T(lab, name)

# 19 · FICHA ROOFTOP INTERNO
S = B.dup(30)
S.T(938, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP INTERNO")
S.T(939, "Ficha de premissas: lazer coberto e copa")
S.P(943, RF(3), (0.5, 0.5)); label_w(S, 944, "CONCEITO SEAZONE · ROOFTOP")
S.P(945, RF(2), (0.5, 0.5)); label_w(S, 946, "COPA · GRANITO PRETO")
S.T(947, "Balcão em madeira voltado para o mar, com banquetas pretas; copa com bancada em granito São Gabriel e armários escuros.")
rows(S, ROWS30, [
    ("Paredes", "Pintura acrílica acetinada cinza médio · cimento queimado na copa (Briefing)"),
    ("Piso", "Porcelanato 90×90 natural/acetinado · rodapé branco de poliestireno"),
    ("Forro", "Forro modular (Briefing)"),
    ("Bancada da copa", "Granito São Gabriel polido · cuba inox · torneira de mesa bica alta"),
    ("Marcenaria", "Balcão e mini market em MDF Carvalho Treviso · armários da copa em MDF Chumbo"),
    ("Mobiliário", "Banquetas e cadeiras pretas, mesas, poltronas (memorial de interiores)"),
    ("Decoração", "Nome do Spot preto com LED no \"O\", frase e letreiro do mini market em preto"),
    ("Copa", "Frigobar + máquina de gelo · mini market com pontos de LED (decisão 19/05/2025)"),
    ("Iluminação", "Spots de embutir pretos na área interna · 2 arandelas na área externa")], PILLS30)

# 20 · FICHA ROOFTOP EXTERNO
S = B.dup(32)
S.T(1007, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP EXTERNO")
S.T(1008, "Ficha de premissas: rooftop externo e piscina")
S.P(1012, RF(7), (0.5, 0.5)); label_w(S, 1013, "CONCEITO · PISCINA")
S.P(1014, RF(6), (0.5, 0.5)); label_w(S, 1015, "PAREDE MOLEDO + LOGO")
S.T(1016, "Parede externa em pedra Moledo com logo CACHOEIRA SPOT e frase \"Seu lugar fora de casa\", banco de concreto com almofadas.")
rows(S, ROWS32, [
    ("Piscina", "Revestimento esverdeado 20×20 (Briefing) · compatibilizar memorial"),
    ("Borda / banco", "Porcelanato que imita madeira 120×20 (Briefing)"),
    ("Guarda-corpo", "Vidro com perfil \"U\" inox (Briefing)"),
    ("Piso externo", "Porcelanato 90×90 antiderrapante (Briefing)"),
    ("Parede Moledo", "Pedra Moledo + banco em estrutura de concreto (LOG) · logo iluminado com ponto elétrico"),
    ("Parede verde", "Jardim vertical com balanço · reforço metálico no teto para o balanço (LOG)"),
    ("Aquecimento", "Aquecimento solar (Briefing) · projeto específico no executivo"),
    ("Ducha", "Ducha de piso + torneira junto ao banheiro PNE (Briefing) · incluir no hidráulico"),
    ("Projeto", "Piscina detalhada no executivo, seguindo este book (conferência 22.0)")], PILLS32)

# 21 · PONTOS DE ATENÇÃO INTERIORES
S = B.dup(24)
S.T(767, "02 · INTERIORES DAS ÁREAS COMUNS · PONTOS DE ATENÇÃO")
S.T(768, "Áreas comuns: pontos de atenção na obra")
S.P(772, RF(6), (0.5, 0.5))
label_w(S, 773, "ROOFTOP · MOLEDO E LOGO")
checklist(S, [
    ("Sem executivo de interiores", "Vale o conceito Seazone e este book; qualquer troca de material precisa de validação registrada."),
    ("Logos iluminados", "Recepção e rooftop: prever fixação e ponto elétrico antes da pedra Moledo."),
    ("Forros", "Gesso acartonado na recepção e forro modular nos corredores e no rooftop (Briefing)."),
    ("Piscina", "Esverdeado 20×20 (Briefing), borda em porcelanato madeira 120×20; aquecimento solar no executivo."),
    ("Balanço e parede verde", "Prever reforço metálico no teto para o balanço; irrigação da parede verde no executivo."),
    ("Marcenaria do rooftop", "Balcão em MDF Carvalho Treviso e armários em MDF Chumbo (estilo Industrial): compatibilizar o memorial.")])

# 22 · DIVISOR 03
S = B.dup(33)
S.P(1048, RF(6), (0.5, 0.5))
S.T(1051, "Materiais de interiores das áreas comuns")
S.geom(1052, y=S.pos(1052)[1] + 0.75)
S.T(1052, "Mapa de materialidade T13–T24: amostras tiradas dos renders Seazone e padrões do catálogo Industrial para compra e aprovação")

# 23-25 · MATERIAIS
NOTE_M = "Amostras extraídas dos renders de interiores Seazone. Referência visual, não substitui amostra física."
for title, note, cards in [
    ("Materiais · paredes (1/3)", NOTE_M, [
        ("T13", "Pedra Moledo", "Parede do logo na recepção e parede externa do rooftop.",
         "Pedra Moledo (LOG de considerações e Briefing) com logo CACHOEIRA SPOT iluminado.",
         "Prever ponto elétrico e fixação do logo antes da pedra.", O("moledo")),
        ("T14", "Cimento queimado areia", "Segunda parede da recepção e copa do rooftop.",
         "Textura cimento queimado areia (Briefing) · frase em letras pretas.",
         "Aplicação por equipe única; aprovar painel-teste no local.", O("cimento")),
        ("T15", "Ripas de madeira", "Hall de entrada da recepção.",
         "Ripas verticais suspensas em MDF Carvalho Treviso (amadeirado do catálogo Industrial).",
         "Fixação detalhada no executivo.", O("ripas")),
        ("T16", "Parede verde", "Área externa do rooftop, parede do balanço.",
         "Parede verde, como na imagem (Briefing: pedra Moledo e parede verde na área externa).",
         "Sistema, irrigação e reforço metálico do balanço detalhados no executivo.", O("verde"))]),
    ("Materiais · rooftop e piscina (2/3)", NOTE_M, [
        ("T17", "Balcão em madeira", "Lazer coberto do rooftop, voltado para o mar.",
         "Balcão em MDF Carvalho Treviso, amadeirado como na imagem, com banquetas pretas.",
         "Aprovar amostra do MDF antes do pedido de marcenaria.", O("bancada")),
        ("T18", "Granito São Gabriel", "Bancada da copa do rooftop.",
         "Granito São Gabriel polido · cuba inox · torneira de mesa bica alta.",
         "Deixar pontos hidráulicos e elétricos conforme o Briefing.", O("saogabriel")),
        ("T19", "Piscina", "Rooftop, com banco e borda.",
         "Revestimento esverdeado 20×20 (Briefing).",
         "A imagem mostra tom acinzentado: seguir o Briefing e aprovar amostra antes do pedido.", O("piscina")),
        ("T20", "Borda e banco amadeirados", "Borda da piscina e banco do rooftop externo.",
         "Porcelanato que imita madeira 120×20 (Briefing).",
         "Peça antiderrapante; conferir caimento e paginação com a imagem.", RF(7), (0.55, 0.78), 0.12)]),
]:
    S = B.dup(34)
    S.T(1060, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
    S.T(1061, title); S.T(1065, note)
    for ids, c in zip(MC, cards):
        card(S, ids, *c)

S = B.dup(37)
S.T(1141, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
S.T(1142, "Materiais · MDFs e metais (3/3)")
S.T(1146, "Padrões Berneck e cores de metal do catálogo Industrial Seazone (rev. 13/10/2026). Referência visual, não substitui amostra física.")
for ids, c in zip(MD, [
        ("T21", "MDF Carvalho Treviso", "Balcão e mini market do rooftop, ripas da recepção.",
         "MDF Berneck Carvalho Treviso (Design), amadeirado claro.",
         "Compatibilizar memorial de interiores (Rícori Nativo); aprovar amostra física antes da compra.", sw("batel_amostra_carvalho_treviso.png")),
        ("T22", "MDF Chumbo", "Armários da copa do rooftop.",
         "MDF Berneck Chumbo (Micro), liso.",
         "Compatibilizar memorial de interiores (Grafito Chess); aprovar amostra física antes da compra.", sw("amostra_chumbo.png")),
        ("T23", "Metal preto RAL 9005", "Pergolado, estruturas de mobiliário, arara, escada e metal da marcenaria.",
         "Metal pintado preto RAL 9005 (catálogo Industrial).",
         "Mesmo acabamento em todas as peças metálicas; aprovar amostra de cor.", sw("batel_amostra_metal_preto.png")),
        ("T24", "Metal grafite RAL 7024", "Segunda cor de metal do catálogo Industrial. Sem uso definido no projeto atual.",
         "Metal pintado grafite RAL 7024 (catálogo Industrial).",
         "Usar só se o projeto pedir segunda cor de metal.", sw("metal_grafite.png"))]):
    card(S, ids, *c)

# 26 · ENCERRAMENTO
S = B.dup(38)
S.P(1168, E["E07"], (0.5, 0.5))
S.T(1170, "CACHOEIRA SPOT")
S.T(1171, "Book de Premissas de Fachada e Interiores das Áreas Comuns")
S.geom(1172, y=S.pos(1172)[1] + 0.55)

finish(B, OUT, "CACHOEIRA SPOT")
