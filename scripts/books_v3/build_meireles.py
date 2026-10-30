"""Meireles Spot no layout v3 (template Campeche). Fonte = book legado da Rachel (pptx) com as regras de 15/10/2026.
python build_meireles.py DADOS saida.pptx   (DADOS = trabalho/v3_dados/meireles)
Estilo Industrial. Trocas aprovadas: MDF Cerrado (Arauco) -> MDF Berneck Carvalho Treviso · metal preto fosco -> RAL 9005.
"""
import glob
import os
import sys

sys.path.insert(0, os.path.expanduser("~/Claude/.claude/skills/book-premissas/scripts/v3"))
from bookv3 import *  # noqa: F401,F403
from pptx.util import Pt  # noqa: E402

D, OUT = sys.argv[1], sys.argv[2]
OLD = os.path.join(REPO, "books", "legado", "MEIRELES_SPOT_Book_de_Premissas_legado.pptx")
src = lambda n: os.path.join(D, "src", n + ".jpg")
sw = lambda n: os.path.join(D, "..", "sw_common", n)
O = lambda k: glob.glob(os.path.join(D, "old", k + "_*"))[0]
E = lambda n: src(f"E{n:02d}")
REC = lambda n: src(f"REC{n:02d}")
RF = lambda n: src(f"ROOF{n:02d}")
E02V3 = src("E02V3")
# amostras do legado (pic id = id do rotulo - 1 ou - 2)
CROMIO, GRANITINA, RIPADO_DIR, MADEIRA_DIR = O("s05_15"), O("s05_19"), O("s05_23"), O("s05_27")
PEDRA_AC, ESQ = O("s06_40"), O("s06_47")
BRISE, PORTAO, CASINHA = O("s07_12"), O("s07_19"), O("s07_33")
DECK, GC = O("s08_26"), O("s08_33")
IND08 = O("s13_12")
MOLEDO, RIPADO_PVC, VERDE, PENDENTES = O("s19_12"), O("s19_19"), O("s19_26"), O("s19_40")
SAOGABRIEL = O("s20_19")
CARVALHO = sw("amostra_carvalho_treviso.png")
PRETO, GRAFITE = sw("amostra_metal_preto_ral9005.png"), sw("amostra_metal_grafite_ral7024.png")

B = Book(TPL, os.path.join(D, "cache"))

# 1 · CAPA
S = B.dup(1)
S.P(16, E(1), (0.5, 0.5))
S.P(17, E(6), (0.45, 0.5))
S.P(18, REC(3), (0.5, 0.5))
S.P(19, RF(10), (0.5, 0.5))
S.T(21, "BOOK DE PREMISSAS DE FACHADA E INTERIORES DAS ÁREAS COMUNS")
S.sh(21).text_frame.paragraphs[0].runs[0].font.size = Pt(8.5)
S.T(22, "Meireles Spot")
S.sh(22).text_frame.paragraphs[0].runs[0].font.size = Pt(36)
S.T(24, "Lançamento  ×  Projeto  ·  Fachada  ·  Interiores das áreas comuns  ·  Estilo Industrial")
S.T(25, "Rua Nunes Valente, 255 · Meireles · Fortaleza/CE", "DOCUMENTO TÉCNICO INTERNO")

# 2 · CONTEXTUALIZAÇÃO
S = B.dup(2)
S.T(36, "Consolidar as premissas de fachada e de interiores das áreas comuns do Meireles Spot a partir da leitura cruzada entre "
        "as imagens de lançamento (REV06), a diretriz de texturas da Arquitetura Seazone, o projeto legal da Projeto Idea, o "
        "Briefing de premissas e o projeto de interiores Seazone: o que a imagem vendeu, o que o projeto traz e como a obra deve executar.")
S.P(37, E(1), (0.5, 0.42), 0.6)
S.P(40, RIPADO_DIR, (0.4, 0.45), 0.6)
S.P(43, RF(10), (0.55, 0.6), 0.45)
S.T(45, "CONFIRA", "o que a obra precisa observar e seguir")
for sid, q in zip([49, 51, 53, 55, 57, 59, 61], [
        "O que as imagens de lançamento mostram na fachada?",
        "Qual é a diretriz de cores e texturas da fachada?",
        "Onde o projeto diverge da imagem vendida?",
        "Como executar o pórtico, as vigas e o pergolado?",
        "Qual é o conceito de interiores das áreas comuns que deve ser seguido?",
        "O que não pode ser alterado sem validação?",
        "O que a obra precisa observar na execução?"]):
    S.T(sid, q)
S.T(62, "Regra do material: premissas validadas pela Seazone. Estilo de interiores: Industrial. Prevalecem a imagem de lançamento "
        "(REV06), o conceito Seazone e o Briefing; o projeto compatibiliza. Fontes (1.28 - [4284] Meireles Spot / 02 - Projetos): "
        "00 - Premissas de Projeto, 08 - Projeto Arquitetônico (imagens 3D, anteprojeto, legal, interiores), 09 - Memorial Descritivo "
        "e diretriz \"MEIRELES SPOT_TEXTURAS\". Interiores das unidades, índices urbanísticos, áreas e custos não fazem parte deste Book.")

# 3 · COMO USAR
S = B.dup(3)
S.P(75, IND08, (0.3, 0.5))
S.T(77, "IND 08")
S.T(78, "Código", "Código da textura (T03) ou da indicação de premissas (IND 08) para achar o item no projeto.")
S.P(87, E(2), (0.5, 0.5), None, 0.8)
S.P(90, REC(3), (0.5, 0.5))
S.P(93, RF(11), (0.5, 0.5))
S.T(89, "Fachada", "imagens de lançamento · paleta · leitura geral · texturas T01–T12 · vistas do topo, da marquise e da base · pontos de atenção")
S.T(92, "Interiores das áreas comuns", "conceito × decisão · recepção e corredores · rooftop interno e externo · pontos de atenção")
S.T(95, "Materiais de interiores das áreas comuns", "texturas T13–T23 para compra e aprovação de amostras, com MDF e metais do estilo Industrial")
S.T(79, "STATUS DOS ITENS")
S.T(81, "premissa validada pela Seazone: seguir na compra e na execução")
S.delete(82, 83, 84, 85)

# 4 · DIVISOR 01
S = B.dup(4)
S.P(101, E(2), (0.5, 0.35))
S.T(105, "Imagens de lançamento  ·  paleta  ·  leitura geral  ·  texturas  ·  vistas com marcação e especificação  ·  pontos de atenção")

# 5 · MOODBOARD FACHADA
S = B.dup(5)
S.P(118, E02V3, (0.5, 0.5)); label_w(S, 119, "PERSPECTIVA")
S.P(120, E(3), (0.5, 0.55)); label_w(S, 121, "MARQUISE")
S.P(122, E(1), (0.5, 0.5)); label_w(S, 123, "FRENTE")
S.P(124, E(5), (0.5, 0.5)); label_w(S, 125, "TORRE")
S.P(126, E(6), (0.45, 0.5)); label_w(S, 127, "BASE E ACESSO")
S.P(128, E(8), (0.5, 0.5)); label_w(S, 129, "VISTA AÉREA")
for sid, (a, b) in zip([132, 134, 136, 138, 140], [
        ("Torre em crômio", "volume cinza texturizado na cor crômio, inclusive o reservatório"),
        ("Concreto ripado", "molduras escuras em torno das janelas e marquise do térreo"),
        ("Pórtico amadeirado", "pórtico do topo, vigas laterais e pergolado do rooftop"),
        ("Base em granitina", "muro cinza com a casinha Seazone iluminada e portão de ripas pretas"),
        ("Identidade", "letreiro MEIRELES SPOT sobre a marquise e logo iluminado no reservatório")]):
    S.T(sid, a, b)

# 6 · PALETA
S = B.dup(6)
for pic, lab, txt, code, a, b, img in [
        (151, 152, 153, "T01", "Crômio texturizado", "torre e reservatório", CROMIO),
        (154, 155, 156, "T02", "Concreto ripado escuro", "molduras e marquise", (RIPADO_DIR, (0.3, 0.4), 0.7)),
        (157, 158, 159, "T03", "Tom amadeirado", "pórtico, vigas e pergolado", MADEIRA_DIR),
        (160, 161, 162, "T04", "Granitina cinza", "muro e base do acesso", GRANITINA),
        (163, 164, 165, "T06", "Esquadrias pretas", "linha 32 (gold), perfil preto", ESQ),
        (166, 167, 168, "T07", "Brise madeira escuro", "ar-condicionado da lateral", BRISE)]:
    S.P(pic, *img) if isinstance(img, tuple) else S.P(pic, img)
    S.T(lab, code); S.T(txt, a, b)
S.T(170, "COMO LER", "Fachada = torre e reservatório em crômio texturizado (T01) + molduras e marquise em concreto ripado escuro (T02) + "
                     "pórtico, vigas e pergolado em tom amadeirado (T03) + muro em granitina cinza (T04) + esquadrias pretas (T06) + "
                     "brise escuro do ar-condicionado (T07). Detalhe de cada material nas texturas T01–T12.")

# 7 · LEITURA GERAL (imagem inteira: quadro ajustado ao aspecto da E07 para caber do logo do topo ao muro)
S = B.dup(7)
S.geom(181, w=6.86)
win = S.P(181, E(7))
label_w(S, 182, "LANÇAMENTO · FRENTE (REV06)")
place_markers(S, 181, win, list(range(183, 193)), [
    (0.585, 0.045), (0.44, 0.122), (0.525, 0.41), (0.58, 0.25), (0.53, 0.66),
    (0.47, 0.698), (0.455, 0.555), (0.52, 0.915), (0.60, 0.955), (0.345, 0.955)])
S.T(193, "A torre vendida (REV06) é um volume cinza texturizado na cor crômio, com molduras de concreto ripado em torno das "
         "janelas, pórtico e vigas em tom amadeirado e esquadrias pretas. No térreo, a marquise de concreto ripado recebe floreira "
         "e o letreiro; o muro em granitina cinza leva a casinha Seazone iluminada. No topo, o reservatório recebe o logo iluminado.")
for sid, (a, b) in zip(range(196, 215, 2), [
        ("Logo no reservatório", "casinha Seazone iluminada no topo"),
        ("Pórtico do topo", "tom amadeirado, como na imagem"),
        ("Vigas laterais", "tom amadeirado, como na imagem"),
        ("Volume do reservatório", "crômio texturizado, plano sem aberturas"),
        ("Fachada da torre", "crômio texturizado (diretriz e memorial)"),
        ("Molduras das janelas", "concreto ripado escuro"),
        ("Esquadrias", "linha 32 (gold), perfil preto"),
        ("Letreiro", "MEIRELES SPOT iluminado sobre a marquise"),
        ("Muro", "granitina cinza, com logo e nome"),
        ("Logo no muro", "casinha branca iluminada")]):
    S.T(sid, a, b)

# 8-10 · TEXTURAS DA FACHADA
NOTE = ("Amostras da diretriz de texturas (nov/2024) e das imagens REV06: referência visual, não substitui amostra física. "
        "Especificação: diretriz, Briefing, Memorial Descritivo e projeto legal.")
for title, cards in [
    ("Texturas da fachada · volumes e revestimentos (1/3)", [
        ("T01", "Crômio texturizado", "Paredes da torre, volumes laterais e reservatório (todas as faces).",
         "Textura/massa acrílica + pintura na cor crômio, texturizada (diretriz nov/2024 e memorial); cinza texturizado da imagem.",
         "Mesma cor e textura na torre e no reservatório; painel-teste aprovado contra a imagem antes da compra.", CROMIO),
        ("T02", "Concreto ripado escuro", "Molduras das janelas, marquise do térreo e laje técnica.",
         "Concreto ripado escuro, como na imagem (diretriz: concreto ripado na laje técnica).",
         "Executivo detalha o sistema (forma ripada ou placa); painel-teste aprovado contra a imagem antes de executar.",
         RIPADO_DIR, (0.4, 0.45), 0.6),
        ("T03", "Pórtico, vigas e pergolado", "Pórtico do topo, vigas da lateral e pergolado do rooftop.",
         "Tom amadeirado, como na imagem REV06 · memorial: aço corten ou madeira · pilares de apoio no pergolado (IND 08).",
         "Seguir a imagem, sem concreto ripado aparente (diretriz anterior à REV06); Projetos detalha material e fixação no executivo.",
         RF(9), (0.925, 0.55), 0.15),
        ("T04", "Granitina cinza", "Muro frontal e base do acesso.",
         "Granitina na cor cinza (diretriz e imagem) com a casinha Seazone e o nome (Briefing: muro com logo e nome).",
         "Visual da imagem + logo e nome; compatibilizar o Briefing (pintura cor Seazone).", GRANITINA)]),
    ("Texturas da fachada · acesso, esquadrias e brises (2/3)", [
        ("T05", "Pedra do acesso", "Lateral do acesso de pedestres, sob a marquise.",
         "Placa clara com veios, como na imagem de lançamento (REV06).",
         "Manter como na imagem; compatibilizar a diretriz de texturas, que pedia a retirada.", PEDRA_AC),
        ("T06", "Esquadrias pretas", "Janelas da torre e esquadrias do térreo e das áreas comuns.",
         "Linha 32 (gold) com perfil preto (Briefing); áreas comuns em preto fosco · legal: alumínio + vidro.",
         "Converge com a imagem: seguir o Briefing no pedido.", ESQ),
        ("T07", "Brise da lateral", "Fechamento das condensadoras de ar-condicionado na lateral.",
         "Brise em madeira, aletas no tom escuro da imagem de lançamento; pode ser substituído por alumínio com o mesmo tom e acabamento.",
         "Preferência madeira tratada para área externa; se a obra optar por alumínio, manter o mesmo tom e acabamento; aprovar amostra contra a imagem.", BRISE),
        ("T08", "Portão de ripas pretas", "Acesso de veículos.",
         "Ripas horizontais pretas (imagem) · portão GA-1 em alumínio, basculante (projeto legal).",
         "Alumínio na cor preta, como na imagem.", PORTAO)]),
    ("Texturas da fachada · identidade e topo (3/3)", [
        ("T09", "Letreiro MEIRELES SPOT", "Sobre a marquise do acesso.",
         "Letras brancas iluminadas, como na imagem · nome do Spot em LED (memorial de interiores e memorial).",
         "Como na imagem; Projetos detalha fixação e ponto elétrico antes do fechamento da fachada.", E(3), (0.62, 0.71), 0.12),
        ("T10", "Casinhas Seazone", "Muro, ao lado da entrada de veículos, e topo do reservatório.",
         "Casinha branca iluminada em grande escala no muro e logo iluminado no reservatório, como nas imagens.",
         "Projetos detalha material, fixação e ponto elétrico antes do fechamento.", CASINHA),
        ("T11", "Floreira da marquise", "Sobre a marquise do térreo.",
         "Vegetação contínua sobre a marquise, como na imagem REV06.",
         "Impermeabilização e irrigação no executivo; floreiras da diretriz fora da REV06 não entram.", E(6), (0.25, 0.5), 0.12),
        ("T12", "Guarda-corpo do rooftop", "Borda do rooftop e da piscina.",
         "Vidro incolor com perfil \"U\" inox (Briefing, área comum) · memorial: aço inox com vidro.",
         "A imagem mostra montantes pretos: seguir o Briefing e compatibilizar o executivo.", GC)])]:
    S = B.dup(8)
    S.T(221, title); S.T(225, NOTE)
    for ids, c in zip(TEX, cards):
        card(S, ids, *c)

# 11 · VISTA DO TOPO (tpl 11)
S = B.dup(11)
S.T(302, "Topo: pórtico, pergolado e rooftop")
win = S.P(306, E(2), (0.45, 0.26), None, 0.8)
label_w(S, 307, "LANÇAMENTO · TOPO (REV06)")
place_markers(S, 306, win, list(range(308, 317)), [
    (0.55, 0.14), (0.34, 0.225), (0.62, 0.36), (0.435, 0.33), (0.313, 0.35), (0.52, 0.225), (0.585, 0.196), (0.357, 0.29)])
S.T(318, "LEITURA DO TOPO", "O pórtico em tom amadeirado envolve o topo da torre e se prolonga no pergolado do rooftop, com pilares "
         "de apoio pedidos pela diretriz (IND 08). Executar no tom amadeirado da imagem; Projetos detalha material e fixação no executivo.")
S.T(320, LEG)
vista_items(S, V11, [
    ("Pórtico", "Tom amadeirado, como na imagem (IND 07)"),
    ("Pergolado", "Tom amadeirado com pilares de apoio (diretriz · IND 08)"),
    ("Vigas laterais", "Tom amadeirado, como na imagem e na diretriz"),
    ("Fachada", "Crômio texturizado (diretriz e memorial)"),
    ("Molduras", "Concreto ripado escuro, como na imagem"),
    ("Esquadrias", "Pretas, linha 32 (gold) com perfil preto (Briefing)"),
    ("Guarda-corpo do rooftop", "Vidro incolor com perfil \"U\" inox (Briefing)"),
    ("Forro das lajes", "Spots embutidos · luminotécnico antes do fechamento")])

# 12 · VISTA DA MARQUISE (tpl 12)
S = B.dup(12)
S.T(354, "Marquise, molduras e lateral")
win = S.P(358, E(3))
label_w(S, 359, "LANÇAMENTO · MARQUISE")
place_markers(S, 358, win, list(range(360, 369)), old_markers(OLD, 10))
S.T(370, LEG)
vista_items(S, V12, [
    ("Pórtico", "Tom amadeirado, como na imagem · material no executivo"),
    ("Molduras", "Concreto ripado escuro, como na imagem (T02)"),
    ("Fachada", "Crômio texturizado (diretriz e memorial)"),
    ("Brise lateral", "Madeira, aletas escuras como na imagem · ou alumínio mesmo tom"),
    ("Esquadrias", "Pretas, linha 32 (gold) com perfil preto (Briefing)"),
    ("Forro das lajes", "Spots embutidos · luminotécnico antes do fechamento"),
    ("Letreiro", "MEIRELES SPOT em letras brancas iluminadas"),
    ("Marquise", "Concreto ripado escuro com floreira, como na imagem"),
    ("Pedra do acesso", "Placa clara com veios, mantida como na imagem")])
S.P(398, RIPADO_DIR, (0.4, 0.45), 0.6); S.P(399, CROMIO); S.P(400, PEDRA_AC)
S.T(401, "T02 · T01 · T05")
S.T(403, "LEITURA DA MARQUISE", "As molduras e a marquise são o principal uso do concreto ripado na imagem. Executar em concreto "
         "ripado escuro, com o sistema detalhado no executivo e painel-teste aprovado contra a imagem antes de executar.")

# 13 · VISTA DA BASE (tpl 13)
S = B.dup(13)
S.T(410, "Base e acesso")
win = S.P(414, E(6))
label_w(S, 415, "LANÇAMENTO · BASE")
pts = old_markers(OLD, 9)
pts[5] = (0.675, 0.555)  # ao lado do letreiro, sem cobrir as letras
pts[2] = (0.19, 0.17)  # abaixo do rotulo da imagem
place_markers(S, 414, win, list(range(416, 426)), pts)
S.T(427, "ATENÇÃO · IDENTIDADE DO ACESSO", "A base é um muro contínuo em granitina cinza com a casinha Seazone em grande escala; a "
         "marquise de concreto ripado recebe floreira e o letreiro iluminado. Projetos detalha fixação, iluminação e pontos elétricos "
         "antes do fechamento da fachada.")
S.T(429, LEG)
vista_items(S, V13, [
    ("Torre", "Crômio texturizado (diretriz e memorial)"),
    ("Molduras", "Concreto ripado escuro, como na imagem"),
    ("Esquadrias", "Pretas, linha 32 (gold) com perfil preto"),
    ("Floreira", "Vegetação contínua · impermeabilização e irrigação"),
    ("Marquise", "Concreto ripado escuro, como na imagem"),
    ("Letreiro", "MEIRELES SPOT em letras brancas iluminadas"),
    ("Logo", "Casinha branca iluminada no muro"),
    ("Muro", "Granitina cinza com logo e nome"),
    ("Portão", "Ripas horizontais pretas · GA-1 em alumínio"),
    ("Acesso", "Porta de pedestres recuada, sob a marquise")])

# 14 · PONTOS DE ATENÇÃO FACHADA
S = B.dup(24)
S.T(767, "01 · FACHADA · PONTOS DE ATENÇÃO")
S.T(768, "Fachada: pontos de atenção na obra")
S.P(772, E(3), (0.5, 0.5))
label_w(S, 773, "LANÇAMENTO · MARQUISE")
checklist(S, [
    ("Pórtico, vigas e pergolado", "Tom amadeirado, como na imagem, com pilares de apoio no pergolado (IND 08); Projetos detalha o material no executivo."),
    ("Concreto ripado", "Molduras, marquise e laje técnica: executivo detalha o sistema; painel-teste aprovado contra a imagem."),
    ("Crômio texturizado", "Torre e reservatório na mesma cor e textura; o executivo traz o quadro de materiais T01–T12."),
    ("Muro e pedra do acesso", "Granitina cinza com casinha e nome (imagem + Briefing); pedra clara do acesso mantida como na imagem."),
    ("Letreiro, casinhas e luz", "Fixação e pontos elétricos antes do fechamento; luminotécnico das lajes entregue por Projetos antes da fachada."),
    ("Brise e floreira", "Brise em madeira escura como na imagem (ou alumínio, mesmo tom); floreira com impermeabilização e irrigação.")])

# 15 · DIVISOR 02
S = B.dup(14)
S.P(465, REC(3), (0.5, 0.5))
S.T(468, "Interiores das áreas comuns")
S.T(469, "Conceito Seazone × decisão  ·  recepção e corredores  ·  rooftop: integração e mini market  ·  piscina  ·  pontos de atenção")

# 16 · CONCEITO × DECISÃO
S = B.dup(15)
S.T(477, "02 · INTERIORES DAS ÁREAS COMUNS")
S.T(478, "Conceito Seazone × decisão registrada")
concept_cards(S, [
    ("Recepção", REC(3), (0.5, 0.5),
     "Parede de pedra com logo Meireles SPOT, ripado amadeirado, jardim vertical, balanço de fibra, espelho meia-lua",
     "Pedras naturais + painel PVC ripado + jardim vertical (Briefing) · reforço na laje para o balanço"),
    ("Rooftop interno", RF(1), (0.5, 0.5),
     "Mesa de madeira, pendentes de fibra, painel ripado com logo, jardim vertical, quadros de cordel",
     "Jardim vertical preservado + painel PVC ripado (Briefing)"),
    ("Copa e mini market", RF(6), (0.6, 0.5),
     "Copa com bancada preta, mini market com estante iluminada",
     "Granito São Gabriel + MDF Berneck Carvalho Treviso · pontos no elétrico"),
    ("Rooftop externo", RF(12), (0.5, 0.5),
     "Pergolado amadeirado, piscina verde, deck, puffs, parede \"Seazone\"",
     "Piscina esverdeada 20×20 · porcelanato madeira 120×20 na piscina e na rampa (Briefing)"),
    ("Projeto de interiores", RF(3), (0.5, 0.5),
     "Imagens, memorial de interiores e moodboard final Seazone",
     "Seguir o conceito Seazone (imagens) até a entrega do executivo de interiores")])
S.T(507, "Imagens de interiores Seazone (10 - Projeto de Interiores / 02 - Imagens) × Briefing, memorial de interiores e moodboard final.")
S.T(509, "ATENÇÃO · REGRA PARA A OBRA", "Vale o conceito das imagens, o moodboard final e o memorial de interiores (acabamentos por "
         "ambiente). Logos, frases, jardim vertical, balanço e mini market entram no elétrico, no hidráulico e no estrutural antes do fechamento.")

# 17 · FICHA RECEPÇÃO
S = B.dup(27)
S.T(848, "02 · INTERIORES DAS ÁREAS COMUNS · RECEPÇÃO E CORREDORES")
S.T(849, "Ficha: recepção e corredores")
S.P(853, REC(3), (0.5, 0.5)); label_w(S, 854, "CONCEITO · PEDRA + BALANÇO")
S.P(855, REC(2), (0.5, 0.5)); label_w(S, 856, "CONCEITO · RIPADO + ESPELHO")
S.T(857, "Parede de pedra com logo Meireles SPOT, ripado amadeirado, jardim vertical, balanço de fibra e banco de madeira de demolição.")
rows(S, ROWS27, [
    ("Parede-destaque", "Pedra Moledo Laguna Branca com logo Meireles SPOT iluminado (memorial de interiores) · ponto elétrico do logo"),
    ("Painel", "Parede de ripado PVC amadeirado (Briefing e memorial de interiores)"),
    ("Jardim", "Parede verde de vegetação preservada (memorial de interiores) · fixação e irrigação no executivo"),
    ("Paredes", "Pintura acrílica acetinada cinza médio (Briefing)"),
    ("Forro", "Forro em gesso acartonado (Briefing)"),
    ("Piso", "Porcelanato 90×90 natural/acetinado, rodapé branco de poliestireno (Briefing)"),
    ("Balanço", "Poltrona balanço de fibra presa na laje, com reforço no estrutural (Briefing)"),
    ("Mobiliário", "Banco de madeira de demolição com metal preto RAL 9005, espelho meia-lua (memorial de interiores)"),
    ("Corredores", "Cinza Urbano + Calopsita Suvinil · arandela sobre a bandeira · portas brancas até o forro (Briefing)")], PILLS27)

# 18 · MOODBOARD ROOFTOP (tpl 31, 11 imagens)
S = B.dup(31)
S.T(977, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP · MOODBOARD")
S.T(978, "Rooftop: integração, mini market e piscina")
for pic, img, foc in zip(range(982, 993),
                         [RF(1), RF(2), RF(4), RF(5), RF(3), RF(6), RF(7), RF(8), RF(9), RF(12), src("ROOF_PLANTA")],
                         [(.5, .5), (.5, .5), (.5, .5), (.5, .5), (.5, .5), (.62, .5), (.5, .5), (.5, .5), (.5, .5), (.5, .5), (.5, .5)]):
    S.P(pic, img, foc)
S.T(993, "CONCEITO", "Inspirado na cultura cearense: quadros de cordel e cerâmica pintada, fibras naturais e madeira; mesa de "
         "madeira com pendentes de fibra, painel ripado com logo, jardim vertical, copa com bancada preta e mini market.",
         "PISCINA", "Revestimento esverdeado 20×20, deck amadeirado, puffs, pergolado amadeirado e parede \"Seazone\".",
         "ESTILO", "Industrial: MDF Berneck Carvalho Treviso e metal preto RAL 9005.")
for pic, lab, name, img in [(994, 995, "Ripado", RIPADO_PVC), (996, 997, "Granito", SAOGABRIEL),
                            (998, 999, "Piscina", DECK), (1000, 1001, "Carvalho", CARVALHO)]:
    S.P(pic, img, (0.75, 0.5) if img == DECK else (0.5, 0.5)); S.T(lab, name)

# 19 · FICHA ROOFTOP INTERNO
S = B.dup(30)
S.T(938, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP INTERNO")
S.T(939, "Ficha de premissas: integração, copa e mini market")
S.P(943, RF(3), (0.5, 0.5)); label_w(S, 944, "CONCEITO · PAINEL RIPADO + LOGO")
S.P(945, RF(6), (0.55, 0.5)); label_w(S, 946, "COPA E MINI MARKET")
S.T(947, "Conceito inspirado na cultura cearense: quadros de cordel e cerâmica pintada, fibras naturais e madeira (moodboard).")
rows(S, ROWS30, [
    ("Paredes", "Pintura acetinada cinza médio + jardim vertical preservado + painel PVC ripado com logo (Briefing)"),
    ("Marcenaria", "Balcão, bancada alta, painel e prateleiras do mini market em MDF Berneck Carvalho Treviso"),
    ("Bancada", "Granito São Gabriel polido na copa (memorial de interiores) · pontos da copa em ilha, saindo do piso (Briefing)"),
    ("Forro", "Forro em gesso acartonado (Briefing)"),
    ("Piso", "Porcelanato 90×90 (Briefing)"),
    ("Iluminação", "Spots pretos de sobrepor + pendentes de fibra marrons sobre mesa e copa · interruptores separados (Briefing)"),
    ("Decoração", "Quadros de cordel e cerâmica pintada, referência ao artesanato cearense (moodboard)"),
    ("Sinalização", "Mini market em LED · nome do Spot com LED no \"O\" (memorial de interiores) · pontos no elétrico"),
    ("Esquadrias", "Preto fosco em toda a área comum (Briefing)")], PILLS30)

# 20 · FICHA ROOFTOP EXTERNO
S = B.dup(32)
S.T(1007, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP EXTERNO")
S.T(1008, "Ficha de premissas: rooftop externo e piscina")
S.P(1012, RF(10), (0.5, 0.5)); label_w(S, 1013, "CONCEITO · PISCINA + DECK")
S.P(1014, RF(11), (0.5, 0.5)); label_w(S, 1015, "PAREDE SEAZONE")
S.T(1016, "Pergolado amadeirado, piscina verde, deck e puffs; letreiro \"Seazone\" iluminado na face do reservatório voltada ao rooftop.")
rows(S, ROWS32, [
    ("Piscina", "Revestimento esverdeado 20×20 (Briefing), verde como na imagem"),
    ("Aquecimento", "Solar para a piscina (Briefing e memorial)"),
    ("Borda e rampa", "Porcelanato que imita madeira 120×20 (Briefing) · compatibilizar memorial (cerâmica atérmica ou São Tomé)"),
    ("Piso do terraço", "Porcelanato 90×90 (Briefing)"),
    ("Pergolado", "Tom amadeirado, como na imagem, com pilares de apoio (diretriz · IND 08)"),
    ("Guarda-corpo", "Vidro incolor com perfil \"U\" inox (Briefing) · compatibilizar com a imagem (montantes pretos)"),
    ("Parede Seazone", "Reservatório em crômio texturizado (diretriz) com letreiro Seazone iluminado · ponto elétrico"),
    ("Mobiliário", "Puffs, como na imagem do conceito Seazone"),
    ("Esquadrias", "Portas de vidro em preto fosco (Briefing)")], PILLS32)

# 21 · PONTOS DE ATENÇÃO INTERIORES
S = B.dup(24)
S.T(767, "02 · INTERIORES DAS ÁREAS COMUNS · PONTOS DE ATENÇÃO")
S.T(768, "Áreas comuns: pontos de atenção na obra")
S.P(772, REC(4), (0.6, 0.5))
label_w(S, 773, "CONCEITO · RECEPÇÃO")
checklist(S, [
    ("Executivo de interiores", "Seguir o conceito Seazone (imagens), o moodboard final e o memorial de interiores até a entrega do executivo."),
    ("Logos, frases e LED", "Logo da recepção, painel do rooftop, Seazone do reservatório e mini market: luminotécnico e pontos antes do fechamento."),
    ("Jardim vertical e balanço", "Prever fixação e irrigação do jardim vertical e reforço na laje para a poltrona balanço (Briefing)."),
    ("Portas das áreas comuns", "Brancas com bandeira até o forro (Briefing); as imagens mostram cinza: seguir o Briefing."),
    ("Piscina e deck", "Esverdeado 20×20 com aquecimento solar; borda e rampa em porcelanato madeira 120×20 (Briefing)."),
    ("Marcenaria e metais", "MDF Berneck Carvalho Treviso e metal preto RAL 9005 (estilo Industrial): compatibilizar o memorial de interiores.")])

# 22 · DIVISOR 03
S = B.dup(33)
S.P(1048, RF(6), (0.55, 0.5))
S.T(1051, "Materiais de interiores das áreas comuns")
S.geom(1052, y=S.pos(1052)[1] + 0.75)
S.T(1052, "Mapa de materialidade T13–T23: amostras tiradas das imagens Seazone e padrões do catálogo Industrial para compra e aprovação")

# 23-24 · MATERIAIS
NOTE_M = "Amostras das imagens de interiores Seazone e especificação do memorial de interiores. Referência visual, não substitui amostra física."
for title, cards in [
    ("Materiais · recepção e rooftop interno (1/3)", [
        ("T13", "Pedra Moledo Laguna Branca", "Parede do logo na recepção.",
         "Pedra Moledo Laguna Branca (memorial de interiores) · Briefing: pedras naturais.",
         "Prever fixação e ponto do logo iluminado antes da pedra.", MOLEDO),
        ("T14", "Ripado PVC amadeirado", "Recepção e rooftop.",
         "Parede de ripado PVC amadeirado (memorial de interiores) · Briefing: painel PVC ripado.",
         "Mesmo padrão na recepção e no rooftop; aprovar amostra antes da compra.", RIPADO_PVC),
        ("T15", "Parede verde", "Recepção e rooftop.",
         "Vegetação preservada (memorial de interiores) · Briefing: jardim vertical preservado.",
         "Prever fixação e irrigação no executivo.", VERDE),
        ("T16", "Granito São Gabriel polido", "Bancada da copa do rooftop.",
         "Granito São Gabriel polido (memorial de interiores).",
         "Pontos da copa em ilha, saindo do piso (Briefing).", SAOGABRIEL)]),
    ("Materiais · rooftop e piscina (2/3)", [
        ("T17", "Pendentes de fibra", "Mesa e copa do rooftop.",
         "Suspensão Jacumã com cúpula e luminária Milhuz, marrons (memorial de interiores).",
         "Interruptores separados (Briefing).", PENDENTES),
        ("T18", "Quadros de cordel", "Paredes do rooftop.",
         "Imagens de cordel e cerâmica pintada, referência ao artesanato cearense (moodboard).",
         "Item de decoração: seguir o moodboard final.", RF(6), (0.635, 0.3), 0.3),
        ("T19", "Piscina esverdeada", "Piscina do rooftop.",
         "Revestimento esverdeado 20×20 (Briefing) · memorial: porcelanato para piscina · aquecimento solar.",
         "Converge com a imagem; aprovar amostra antes do pedido.", RF(13), (0.42, 0.75), 0.18),
        ("T20", "Deck amadeirado", "Entorno da piscina e rampa.",
         "Porcelanato que imita madeira 120×20, só na piscina e na rampa (Briefing), visual de deck da imagem.",
         "Peça antiderrapante; conferir paginação com a imagem.", RF(10), (0.42, 0.9), 0.16)])]:
    S = B.dup(34)
    S.T(1060, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
    S.T(1061, title); S.T(1065, NOTE_M)
    for ids, c in zip(MC, cards):
        card(S, ids, *c)

# 25 · MDFs E METAIS (3 cards)
S = B.dup(37)
S.T(1141, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
S.T(1142, "Materiais · MDF e metais (3/3)")
S.T(1146, "Padrão Berneck e cores de metal do catálogo Industrial Seazone (rev. 13/10/2026). Referência visual, não substitui amostra física.")
for ids, c in zip(MD, [
        ("T21", "MDF Carvalho Treviso", "Copa e mini market do rooftop (balcão, bancada alta, painel e prateleiras) e demais marcenarias.",
         "MDF Berneck Carvalho Treviso (Design), amadeirado claro, com metal preto RAL 9005.",
         "Compatibilizar memorial de interiores (Cerrado); aprovar amostra física antes da compra.", CARVALHO),
        ("T22", "Metal preto RAL 9005", "Banco da recepção, estruturas de mobiliário e metais da marcenaria.",
         "Metal pintado preto RAL 9005 (catálogo Industrial).",
         "Compatibilizar memorial de interiores (metal preto fosco); mesmo acabamento em todas as peças; aprovar amostra de cor.", PRETO),
        ("T23", "Metal grafite RAL 7024", "Segunda cor de metal do catálogo Industrial. Sem uso definido no projeto atual.",
         "Metal pintado grafite RAL 7024 (catálogo Industrial).",
         "Usar só se o projeto pedir segunda cor de metal.", GRAFITE)]):
    card(S, ids, *c)
S.delete(1159, 1160, 1161, 1162)
for sid in (1155, 1156, 1157, 1158):  # terceiro card centralizado
    S.geom(sid, x=S.pos(sid)[0] + 3.16)

# 26 · ENCERRAMENTO
S = B.dup(38)
S.P(1168, E(2), (0.5, 0.4))
S.T(1170, "MEIRELES SPOT")
S.T(1171, "Book de Premissas de Fachada e Interiores das Áreas Comuns")
S.geom(1172, y=S.pos(1172)[1] + 0.55)
S.T(1172, "Lançamento × Projeto  ·  Seazone")

finish(B, OUT, "MEIRELES SPOT")
