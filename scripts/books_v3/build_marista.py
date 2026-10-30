"""Marista 144 Spot (Goiania/GO) no layout v3 (template Campeche). Conteudo = book legado da Rachel (pptx) com as
regras do v3 aplicadas (tudo DEFINIDO, estilo Industrial, MDFs Solanum / Cinza Cristal, metal preto RAL 9005).
python build_marista.py DADOS saida.pptx   (DADOS = trabalho/v3_dados/marista)
"""
import os
import sys

sys.path.insert(0, os.path.expanduser("~/Claude/.claude/skills/book-premissas/scripts/v3"))
from bookv3 import *  # noqa: F401,F403
from pptx.util import Pt  # noqa: E402

D, OUT = sys.argv[1], sys.argv[2]
OLD = os.path.join(REPO, "books", "legado", "MARISTA_144_SPOT_Book_de_Premissas_legado.pptx")
src = lambda n: os.path.join(D, "src", n + ".jpg")
swc = lambda n: os.path.join(D, "..", "sw_common", n)
EX = {k: src(k) for k in ["E01", "E02", "E03", "E03N", "E04", "E06", "E07", "E08", "E09", "E10", "E11", "E11N"]}
REC = lambda n: src(f"REC0{n}")
RF = lambda n: src(f"ROOF0{n}")
RECP, RFP = src("REC_PLANTA"), src("ROOF_PLANTA")
OLDI = {"cinza": "s05_12_b15bb44b.jpg", "branco": "s05_19_eb7aeb50.jpg", "ripado": "s05_26_bbc350ca.jpg",
        "esq": "s05_33_6dfa15a4.jpg", "aleta": "s05_40_a1403c01.jpg", "flor": "s05_47_3ce7eece.jpg",
        "portico": "s06_12_fe3577d7.jpg", "gc": "s06_19_e18223e2.jpg", "spots": "s06_26_0c1f2252.jpg",
        "marquise": "s06_33_f2df6ff4.jpg", "casinha": "s06_40_ea1d4953.jpg", "seazone": "s06_47_cb2b21fd.jpg",
        "cimento": "s18_12_0e46be68.jpg", "savana": "s18_19_0b6307a3.jpg", "jardim": "s18_26_7712986c.jpg",
        "brick": "s18_33_d9c8c77a.jpg", "pvc": "s18_40_7f1586ce.jpg", "tinta": "s18_47_4aae554f.jpg",
        "saogabriel": "s19_26_36e8501a.jpg", "porcelanato": "s19_33_d184fdf0.jpg"}
O = lambda k: os.path.join(D, "old", OLDI[k])

B = Book(TPL, os.path.join(D, "cache"))

# ================= 1 · CAPA =================
S = B.dup(1)
S.P(16, EX["E02"], (0.5, 0.5))
S.P(17, EX["E07"], (0.55, 0.55))
S.P(18, REC(2), (0.5, 0.5))
S.P(19, RF(7), (0.5, 0.5))
S.T(21, "BOOK DE PREMISSAS DE FACHADA E INTERIORES DAS ÁREAS COMUNS")
S.sh(21).text_frame.paragraphs[0].runs[0].font.size = Pt(8.5)
S.T(22, "Marista 144 Spot")
S.sh(22).text_frame.paragraphs[0].runs[0].font.size = Pt(36)
S.T(24, "Lançamento  ×  Pré-executivo  ·  Fachada  ·  Interiores das áreas comuns  ·  Estilo Industrial")
S.T(25, "Rua 144 · Setor Marista · Goiânia/GO", "DOCUMENTO TÉCNICO INTERNO")

# ================= 2 · CONTEXTUALIZAÇÃO =================
S = B.dup(2)
S.T(36, "Consolidar as premissas de fachada e de interiores das áreas comuns do Marista 144 Spot a partir da leitura cruzada "
        "entre as imagens de lançamento, o pré-executivo da Norden Arquitetura (fachadas 301 a 303, R03), o Briefing de "
        "premissas e os projetos de interiores (Seazone e ACÁ): o que a imagem vendeu, o que o projeto especifica, onde o "
        "projeto precisa ser ajustado e o que a obra deve seguir na compra e na execução.")
S.P(37, EX["E09"], (0.5, 0.5), 0.6)
S.P(40, O("ripado"))
S.P(43, EX["E07"], (0.6, 0.5), 0.35)
S.T(45, "CONFIRA", "o que a obra precisa observar e seguir")
for sid, q in zip([49, 51, 53, 55, 57, 59, 61], [
        "O que as imagens de lançamento mostram na fachada?",
        "Qual textura e cor o pré-executivo define para cada volume?",
        "Qual é o material do ripado, das aletas e do pórtico?",
        "Como são a marquise, a casinha, os letreiros e as floreiras?",
        "Qual é o conceito de interiores das áreas comuns que deve ser seguido?",
        "O que não pode ser alterado sem validação?",
        "O que a obra precisa observar na execução?"]):
    S.T(sid, q)
S.T(62, "Regra do material: premissas validadas pela Seazone. Estilo de interiores: Industrial. Prevalecem a imagem de "
        "lançamento, o conceito Seazone e o Briefing; o pré-executivo e o executivo compatibilizam. Fontes (1.25 - [3352] "
        "Marista 144 Spot / 02 - Projetos): 08 - Projeto Arquitetônico (imagens 3D, pré-executivo Norden R03, interiores Seazone "
        "e ACÁ), 09 - Memorial Descritivo, 00 - Premissas de Projeto e Conferência de Projetos Executivos. "
        "Interiores das unidades, índices urbanísticos, áreas e custos não fazem parte deste Book.")

# ================= 3 · COMO USAR =================
S = B.dup(3)
S.P(75, O("ripado"))
S.T(77, "T03")
S.T(78, "Código", "Código do book (T01 a T24) para achar o item nas páginas de texturas e materiais.")
S.P(87, EX["E09"], (0.5, 0.5), None, 0.8)
S.P(90, REC(2), (0.5, 0.5))
S.P(93, RF(7), (0.5, 0.5))
S.T(89, "Fachada", "imagens de lançamento · paleta · leitura geral · texturas T01–T12 · vistas esquina, detalhe e acesso · pontos de atenção")
S.T(92, "Interiores das áreas comuns", "conceito × decisão · recepção · rooftop interno (copa e mini market) · rooftop externo e piscina · pontos de atenção")
S.T(95, "Materiais de interiores das áreas comuns", "texturas T13–T24 para compra e aprovação de amostras, com MDFs e metais do estilo Industrial")
S.T(79, "STATUS DOS ITENS")
S.T(81, "premissa validada pela Seazone: seguir na compra e na execução")
S.delete(82, 83, 84, 85)

# ================= 4 · DIVISOR 01 =================
S = B.dup(4)
S.P(101, EX["E09"], (0.5, 0.5))
S.T(105, "Imagens de lançamento  ·  paleta  ·  leitura geral  ·  texturas  ·  vistas com marcação e especificação  ·  pontos de atenção")

# ================= 5 · MOODBOARD FACHADA =================
S = B.dup(5)
S.P(118, EX["E11"], (0.5, 0.5)); label_w(S, 119, "11 · FRONTAL")
S.P(120, EX["E07"], (0.5, 0.5)); label_w(S, 121, "07 · ACESSO")
S.P(122, EX["E02"], (0.5, 0.5)); label_w(S, 123, "02")
S.P(124, EX["E03N"], (0.5, 0.5)); label_w(S, 125, "03 · NOITE")
S.P(126, EX["E04"], (0.5, 0.45)); label_w(S, 127, "04 · DETALHE")
S.P(128, EX["E09"], (0.5, 0.5)); label_w(S, 129, "09 · ESQUINA")
for sid, (a, b) in zip([132, 134, 136, 138, 140], [
        ("Cinza + branco", "paredes em textura cinza Munique entre as janelas e bordas de laje em Branco Gelo, desenhando cada pavimento"),
        ("Ripado amadeirado", "painéis verticais em \"L\" ao lado das janelas e fechamento das sacadas, em alumínio amadeirado"),
        ("Aletas grafite", "duas faixas verticais escuras do térreo ao pórtico, marcando a fachada frontal"),
        ("Base e acesso", "marquise escura com jardineira e letreiro MARISTA 144 SPOT, casinha escultórica, cobogó e letreiro seazone"),
        ("Coroamento", "pórtico claro e pergolado metálico preto sobre o rooftop, floreiras nas quinas e guarda-corpo de vidro")]):
    S.T(sid, a, b)

# ================= 6 · PALETA =================
S = B.dup(6)
for pic, lab, txt, code, a, b, img, foc, hf in [
        (151, 152, 153, "T01", "Textura cinza Munique", "paredes das sacadas e entre janelas", EX["E04"], (0.56, 0.30), 0.07),
        (154, 155, 156, "T02", "Textura Branco Gelo", "bordas de laje e volumes claros", EX["E04"], (0.81, 0.553), 0.06),
        (157, 158, 159, "T03", "Brise alumínio amadeirado", "painéis em \"L\" e sacadas", O("ripado"), (0.5, 0.5), None),
        (160, 161, 162, "T04", "Esquadrias preto fosco", "janelas e porta-janelas", O("esq"), (0.5, 0.5), None),
        (163, 164, 165, "T05", "Aletas grafite", "faixas verticais da frente", O("aleta"), (0.5, 0.5), None),
        (166, 167, 168, "T07", "Pergolado preto RAL 9005", "pórtico do rooftop", EX["E11"], (0.475, 0.075), 0.11)]:
    S.P(pic, img, foc, hf); S.T(lab, code); S.T(txt, a, b)
S.T(170, "COMO LER", "Fachada = textura cinza Munique (T01) + bordas e volumes em Branco Gelo (T02) + brise vertical em alumínio amadeirado (T03) "
                     "+ esquadrias pretas (T04) + aletas grafite (T05) + pórtico com pergolado metálico preto (T07). "
                     "Detalhe de cada material nas páginas de texturas T01–T12.")

# ================= 7 · LEITURA GERAL =================
S = B.dup(7)
win = S.P(181, EX["E11"])
label_w(S, 182, "LANÇAMENTO · 11 (FRONTAL)")
PTS7 = [(0.475, 0.071), (0.4425, 0.40), (0.55, 0.268), (0.48, 0.50), (0.39, 0.571),
        (0.50, 0.696), (0.565, 0.624), (0.385, 0.339), (0.53, 0.838), (0.585, 0.883)]
place_markers(S, 181, win, list(range(183, 193)), PTS7)
S.T(193, "A fachada é lida em três camadas: paredes em textura cinza Munique com bordas de laje em Branco Gelo e spots; painéis "
         "ripados amadeirados em \"L\" ao lado das janelas e no fechamento das sacadas; e duas aletas grafite que sobem do térreo "
         "ao pórtico com pergolado do rooftop. Na base, marquise com jardineira e letreiro MARISTA 144 SPOT, casinha escultórica, "
         "cobogó e letreiro seazone iluminado.")
for sid, (a, b) in zip(range(196, 215, 2), [
        ("Pórtico e pergolado do rooftop", "pórtico claro + pergolado metálico preto RAL 9005, como na imagem"),
        ("Aletas verticais grafite", "duas faixas escuras do térreo ao pórtico"),
        ("Ripado vertical amadeirado", "painéis em \"L\" ao lado das janelas · alumínio amadeirado"),
        ("Paredes em textura cinza", "Textucril Leinertex cor Munique (pré-executivo)"),
        ("Bordas brancas com spots", "lajes em balanço em Branco Gelo com iluminação embutida"),
        ("Esquadrias pretas", "linha 32 (gold), caixilho quadriculado preto fosco, vidro incolor"),
        ("Guarda-corpo ripado", "fechamento amadeirado das sacadas, como na imagem"),
        ("Floreiras nas bordas", "vegetação nas quinas do volume · paisagismo PE R01"),
        ("Marquise com letreiro", "MARISTA 144 SPOT em letras brancas + jardineira"),
        ("Logo-casinha e seazone", "escultura no acesso e letreiro iluminado, como na imagem")]):
    S.T(sid, a, b)

# ================= 8-10 · TEXTURAS DA FACHADA =================
NOTE = ("Amostras extraídas das imagens de lançamento: referência visual, não substitui amostra física. "
        "Especificação: pré-executivo Norden (quadro de materiais das fachadas 301 a 303, R03), Premissas, Memorial e Conferência.")
for title, cards in [
    ("Texturas da fachada · cores e ripado (1/3)", [
        ("T01", "Textura cinza Munique", "Paredes das sacadas e entre as janelas, nas quatro fachadas.",
         "Pintura texturizada, revestimento acrílico Textucril Leinertex, cor Munique (pré-executivo); tom cinza como na imagem.",
         "Painel-teste contra a imagem antes da compra; aplicação por equipe única.", EX["E04"], (0.56, 0.30), 0.07),
        ("T02", "Textura Branco Gelo", "Bordas das lajes e volumes laterais claros.",
         "Pintura texturizada Textucril Leinertex, cor Branco Gelo (pré-executivo). Converge com a imagem.",
         "Mesma cor e lote em todas as faces; painel-teste junto com a Munique.", EX["E04"], (0.81, 0.553), 0.06),
        ("T03", "Ripado vertical amadeirado", "Painéis em \"L\" ao lado das janelas e fechamento das sacadas.",
         "Brise em alumínio com acabamento amadeirado, ripas verticais tom madeira como na imagem · memorial: madeira ou brises móveis.",
         "Pré-executivo sem material e memorial (madeira): executar em alumínio, acabamento amadeirado no tom da imagem; aprovar amostra antes da compra.", O("ripado")),
        ("T04", "Esquadrias pretas", "Janelas e porta-janelas de todos os pavimentos e áreas comuns.",
         "Linha 32 (gold), PVC ou alumínio, acabamento preto fosco como na imagem (Briefing: áreas comuns preto fosco) · vidro incolor.",
         "Projeto de esquadrias (Béni): compatibilizar linha e cor antes do pedido; contramarco antes do reboco.", O("esq"))]),
    ("Texturas da fachada · aletas, floreiras e pórtico (2/3)", [
        ("T05", "Aletas verticais grafite", "Duas faixas da fachada frontal, do térreo ao pórtico.",
         "Faixas grafite escuro, como na imagem · referência de cor: grafite RAL 7024 (catálogo Industrial).",
         "Pré-executivo não traz cor própria: completar o quadro de materiais com o grafite; painel-teste contra a imagem.", O("aleta")),
        ("T06", "Floreiras nas bordas", "Quinas das lajes do volume, em todos os pavimentos.",
         "Floreiras brancas com vegetação pendente, como na imagem · espécies e irrigação conforme paisagismo PE R01.",
         "Impermeabilização com pintura antirraiz e teste de estanqueidade antes do plantio.", O("flor")),
        ("T07", "Pórtico e pergolado", "Topo do edifício, sobre o rooftop.",
         "Pórtico em moldura clara (Branco Gelo) e pergolado em estrutura metálica preta RAL 9005, como na imagem (estilo Industrial).",
         "Projeto metálico próprio detalha; compatibilizar memorial (aço corten ou madeira) antes da fabricação.", EX["E11"], (0.475, 0.075), 0.11),
        ("T08", "Guarda-corpo ripado", "Sacadas dos pavimentos-tipo.",
         "Fechamento ripado em alumínio amadeirado como na imagem (mesmo sistema de T03) sobre guarda-corpo em alumínio anodizado com montante aparente (Briefing).",
         "Detalhar perfil, altura e fixação no executivo; compatibilizar o pré-executivo (sem material).", O("gc"))]),
    ("Texturas da fachada · luz e identidade (3/3)", [
        ("T09", "Spots nas bordas", "Face inferior das lajes em balanço, em todos os pavimentos.",
         "Spots embutidos em todas as bordas, como na imagem.",
         "Luminotécnico: Projetos entrega antes do fechamento de fachada; prever os pontos no elétrico.", EX["E04"], (0.545, 0.12), 0.08),
        ("T10", "Marquise com letreiro", "Acesso principal do térreo.",
         "Marquise em concreto escuro com jardineira e letras brancas MARISTA 144 SPOT, como na imagem (memorial: nome na fachada interna).",
         "Executar como na imagem; Projetos detalha estrutura, fixação das letras e ponto elétrico.", EX["E07"], (0.62, 0.47), 0.25),
        ("T11", "Logo-casinha escultórica", "Acesso de veículos, junto ao cobogó.",
         "Casinha branca em grande escala junto ao painel de cobogó do acesso, como na imagem.",
         "Projetos detalha estrutura, fixação e acabamento; executar como na imagem.", O("casinha")),
        ("T12", "Letreiro seazone iluminado", "Parede escura do térreo, ao lado do acesso.",
         "Letras brancas iluminadas sobre parede escura, como na imagem · muro com logo e nome Seazone (Briefing).",
         "Incluir letreiros e logo no elétrico; Projetos detalha fixação antes do fechamento da fachada.", O("seazone"))])]:
    S = B.dup(8)
    S.T(221, title); S.T(225, NOTE)
    for ids, c in zip(TEX, cards):
        card(S, ids, *c)

# ================= 11 · VISTA DETALHE (tpl 11) =================
S = B.dup(11)
S.T(302, "Detalhe: ripado, texturas e esquadrias")
win = S.P(306, EX["E04"], (0.5, 0.45))
label_w(S, 307, "LANÇAMENTO · 04 (DETALHE)")
PTS11 = [(0.45, 0.35), (0.55, 0.29), (0.81, 0.553), (0.926, 0.44), (0.24, 0.505), (0.56, 0.505), (0.66, 0.40), (0.06, 0.40)]
place_markers(S, 306, win, list(range(308, 317)), PTS11)
S.T(318, "LEITURA DO DETALHE", "Cada pavimento é desenhado pela borda de laje em Branco Gelo; entre elas, paredes em textura "
         "cinza Munique, esquadrias pretas e painéis ripados amadeirados em \"L\" que viram fechamento das sacadas. As aletas "
         "grafite e as floreiras nas quinas completam a composição.")
S.T(320, LEG)
vista_items(S, V11, [
    ("Ripado vertical", "Brise em alumínio amadeirado, ripas verticais como na imagem"),
    ("Textura cinza", "Textucril Leinertex cor Munique (pré-executivo)"),
    ("Borda de laje", "Textucril Leinertex cor Branco Gelo, com spots embutidos"),
    ("Esquadrias", "Linha 32 (gold), preto fosco, vidro incolor"),
    ("Floreiras", "Floreiras brancas nas quinas · paisagismo PE R01"),
    ("Guarda-corpo ripado", "Alumínio amadeirado + anodizado, montante aparente"),
    ("Aleta grafite", "Faixa vertical grafite escuro, como na imagem"),
    ("Volume lateral claro", "Branco Gelo (Textucril Leinertex)")])

# ================= 12 · VISTA ESQUINA (tpl 12) =================
S = B.dup(12)
S.T(354, "Esquina e volume lateral")
win = S.P(358, EX["E02"])
label_w(S, 359, "LANÇAMENTO · 02")
place_markers(S, 358, win, list(range(360, 369)), old_markers(OLD, 7))
S.T(370, LEG)
vista_items(S, V12, [
    ("Pórtico", "Moldura clara no topo (Branco Gelo), como na imagem · projeto metálico próprio"),
    ("Volume lateral", "Branco Gelo (Textucril Leinertex)"),
    ("Ripado", "Alumínio amadeirado, como na imagem; compatibilizar o pré-executivo"),
    ("Aletas", "Grafite escuro, como na imagem; completar o quadro de materiais"),
    ("Esquadrias", "Linha 32 (gold), preto fosco (Briefing e imagem)"),
    ("Floreiras", "Vegetação nas quinas · paisagismo PE R01"),
    ("Pergolado", "Metal preto RAL 9005, escuro como na imagem; compatibilizar memorial"),
    ("Térreo", "Marquise, letreiros e casinha como na imagem; Projetos detalha")])
S.P(398, EX["E04"], (0.56, 0.30), 0.07); S.P(399, EX["E04"], (0.81, 0.553), 0.06); S.P(400, O("ripado"))
S.T(401, "T01 · T02 · T03")
S.T(403, "NO PRÉ-EXECUTIVO", "O quadro de materiais das fachadas 301 a 303 (Norden R03) traz só Munique e Branco Gelo. Ripado, "
         "aletas, pórtico, pergolado e guarda-corpos seguem este book (imagem + Briefing) e entram no executivo de arquitetura.")

# ================= 13 · VISTA TÉRREO (tpl 13) =================
S = B.dup(13)
S.T(410, "Térreo: acesso, marquise e identidade")
win = S.P(414, EX["E07"])
label_w(S, 415, "LANÇAMENTO · 07 (ACESSO)")
place_markers(S, 414, win, list(range(416, 426)), old_markers(OLD, 8))
S.T(427, "ATENÇÃO · IDENTIDADE DO ACESSO", "Marquise com jardineira e letreiro, casinha escultórica, cobogó e letreiro seazone "
         "iluminado são a \"cara\" do empreendimento no lançamento: executar como na imagem. Nenhum deles está no pré-executivo: "
         "Projetos detalha estrutura, fixação e pontos elétricos antes do fechamento de fachada.")
S.T(429, LEG)
vista_items(S, V13, [
    ("Letreiro", "MARISTA 144 SPOT em letras brancas iluminadas, como na imagem"),
    ("Marquise", "Concreto escuro, como na imagem · Projetos detalha estrutura"),
    ("Jardineira", "Vegetação sobre a marquise · impermeabilização antirraiz"),
    ("Casinha", "Escultura branca no acesso, como na imagem · Projetos detalha"),
    ("Cobogó", "Painel vazado junto ao acesso, como na imagem"),
    ("seazone", "Letreiro iluminado na parede escura · ponto elétrico"),
    ("Aletas", "Grafite escuro até o térreo"),
    ("Ripado", "Alumínio amadeirado nos pavimentos, como na imagem")])

# ================= 14 · PONTOS DE ATENÇÃO FACHADA =================
S = B.dup(24)
S.T(767, "01 · FACHADA · PONTOS DE ATENÇÃO")
S.T(768, "Fachada: pontos de atenção na obra")
S.P(772, EX["E03N"], (0.5, 0.5))
label_w(S, 773, "LANÇAMENTO · 03 (NOITE)")
checklist(S, [
    ("Ripado e brises", "Alumínio com acabamento amadeirado no tom da imagem: compatibilizar o pré-executivo (sem material) e o memorial (madeira); projeto de brises detalha a fixação."),
    ("Esquadrias", "Linha 32 (gold) preto fosco, como na imagem: compatibilizar o projeto de esquadrias (Béni) antes do pedido."),
    ("Guarda-corpo das sacadas", "Ripado em alumínio amadeirado sobre alumínio anodizado com montante aparente (Briefing); detalhar no executivo."),
    ("Aletas, pórtico e pergolado", "Aletas grafite; pórtico claro; pergolado metálico preto RAL 9005: completar o quadro de materiais e compatibilizar o memorial (corten/madeira)."),
    ("Identidade e luz", "Marquise, letreiros, casinha, cobogó e spots nas bordas como na imagem; luminotécnico e pontos elétricos antes do fechamento de fachada."),
    ("Executivo e texturas", "Emitir o executivo com quadro de materiais completo e prancha 304; painel-teste de Munique e Branco Gelo contra a imagem antes da compra.")])

# ================= 15 · DIVISOR 02 =================
S = B.dup(14)
S.P(465, REC(2), (0.5, 0.5))
S.T(468, "Interiores das áreas comuns")
S.T(469, "Conceito Seazone × decisão  ·  recepção  ·  rooftop interno: copa e mini market  ·  rooftop externo e piscina  ·  pontos de atenção")

# ================= 16 · CONCEITO × DECISÃO =================
S = B.dup(15)
S.T(477, "02 · INTERIORES DAS ÁREAS COMUNS")
S.T(478, "Conceito Seazone × decisão registrada")
concept_cards(S, [
    ("Recepção", REC(2), (0.5, 0.5),
     "Painel de madeira com frisos de LED e letreiro Marista 144 SPOT, jardim vertical, cobogó, cimento queimado, sofá curvo",
     "Seguir o conceito (imagens) com os materiais do projeto básico ACÁ; painel em MDF Berneck Solanum"),
    ("Rooftop interno", RF(1), (0.5, 0.5),
     "Mesa oval, painel ripado com letreiro, tijolinho branco na copa, bancada de granito São Gabriel, mini market",
     "Seguir o conceito; marcenaria em MDF Solanum e Cinza Cristal com metal preto RAL 9005"),
    ("Rooftop externo", RF(5), (0.5, 0.5),
     "Deck amadeirado, banco com floreira, logo-casinha preto na parede, mesas e espreguiçadeiras",
     "Deck em porcelanato que imita madeira 120×20 (Briefing) · logo-casinha como na imagem"),
    ("Piscina", RF(7), (0.5, 0.5),
     "Piscina de borda infinita com revestimento esverdeado e degraus revestidos",
     "Esverdeado 20×20 (Briefing e imagem); projeto próprio da piscina detalha no executivo"),
    ("Executivo de interiores", RECP, (0.5, 0.5),
     "Projeto básico ACÁ no Drive (recepção e rooftop), com materiais e layout do conceito Seazone",
     "Executivo de interiores ACÁ: seguir o conceito Seazone (imagens) e este book até a entrega do executivo")])
S.T(507, "Imagens de interiores (versão atualizada após a alteração do projeto legal) × Conferência de Projetos Executivos "
         "(aba Interiores), Briefing e projeto básico ACÁ.")
S.T(509, "ATENÇÃO · REGRA PARA A OBRA", "Vale o conceito das imagens atualizadas (10 - Projeto de Interiores / Seazone / 02 - Imagens / "
         "ALTERAÇÃO 05-2026) e os materiais do projeto básico ACÁ e do moodboard Seazone. Letreiros, logo iluminado, frisos de LED "
         "e pontos do mini market entram no elétrico antes do fechamento.")

# ================= 17 · FICHA RECEPÇÃO =================
S = B.dup(27)
S.T(848, "02 · INTERIORES DAS ÁREAS COMUNS · RECEPÇÃO")
S.T(849, "Ficha de premissas: recepção e hall")
S.P(853, REC(2), (0.5, 0.5)); label_w(S, 854, "CONCEITO · RECEPÇÃO")
S.P(855, REC(3), (0.5, 0.5)); label_w(S, 856, "CONCEITO · JARDIM VERTICAL")
S.T(857, "Painel de madeira com frisos de LED e letreiro, cobogó com vegetação, jardim vertical e espelho orgânico. Indicação 02: "
         "porta embutida no painel.")
rows(S, ROWS27, [
    ("Paredes", "Cimento queimado Av. Expressa Suvinil (ACÁ e Briefing) · demais paredes Cinza Urbano Suvinil (Briefing); compatibilizar ACÁ (Nuvem de Chuva)"),
    ("Painel", "MDF Berneck Solanum com perfis em metal preto RAL 9005 e frisos de LED, como na imagem (ACÁ: Savana + metalon preto)"),
    ("Destaques", "Jardim vertical preservado (ACÁ e Briefing) com o visual das imagens · cobogó com vegetação (imagem)"),
    ("Forro", "Gesso com pé-direito livre (Briefing); compatibilizar as imagens (forro modular)"),
    ("Piso", "Porcelanato 90×90 natural/acetinado · rodapé branco de poliestireno (Briefing)"),
    ("Identidade", "Letreiro Marista 144 SPOT no painel e logo seazone iluminado (Briefing e imagem) · pontos no elétrico"),
    ("Iluminação", "2 pendentes + spots de sobrepor, luz neutra 4000 K (Briefing) · frisos de LED do painel"),
    ("Mobiliário", "Sofá curvo, poltronas com estrutura preta, banco de madeira, tapetes redondos (imagem)"),
    ("Portas", "Brancas com bandeira até o forro (Briefing) · porta do painel no mesmo MDF Solanum (Indicação 02) · lavabo na tinta da recepção")], PILLS27)

# ================= 18 · MOODBOARD ROOFTOP (tpl 28) =================
S = B.dup(28)
S.T(888, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP · MOODBOARD")
S.T(889, "Rooftop: integração, copa, mini market, deck e piscina")
for pic, img in zip(range(893, 902), [RF(1), RF(3), RF(2), RF(4), RF(5), RF(6), RF(7), RF(8), RF(9)]):
    S.P(pic, img)
S.T(902, "CONCEITO", "Mesa oval com cadeiras, painel ripado amadeirado com letreiro iluminado, copa com tijolinho branco e bancada "
         "escura em granito São Gabriel, mini market com fita de LED e logo iluminado.",
         "EXTERNO", "Deck amadeirado, banco com floreira e logo-casinha preto na parede; piscina de borda infinita esverdeada com "
         "degraus revestidos e guarda-corpo de vidro.",
         "ESTILO", "Industrial: MDF Berneck Solanum e Cinza Cristal com metal preto RAL 9005.")
for pic, lab, name, img, foc, hf in [(903, 904, "Tijolinho", RF(1), (0.85, 0.35), 0.18), (905, 906, "Ripado", RF(1), (0.05, 0.66), 0.08),
                                     (907, 908, "Piscina", RF(7), (0.6, 0.72), 0.18), (909, 910, "Deck", RF(5), (0.45, 0.88), 0.14)]:
    S.P(pic, img, foc, hf); S.T(lab, name)

# ================= 19 · FICHA ROOFTOP INTERNO =================
S = B.dup(30)
S.T(938, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP INTERNO")
S.T(939, "Ficha de premissas: integração, copa e mini market")
S.P(943, RF(1), (0.5, 0.5)); label_w(S, 944, "CONCEITO SEAZONE · COPA")
S.P(945, RF(4), (0.5, 0.5)); label_w(S, 946, "INTEGRAÇÃO · MESA OVAL")
S.T(947, "Mesa oval, painel ripado com letreiro iluminado, tijolinho branco e bancada escura na copa. Indicação 03: portas do rooftop na cor da parede.")
rows(S, ROWS30, [
    ("Paredes", "Cinza Urbano + Calopsita Suvinil (Briefing); compatibilizar ACÁ (Nuvem de Chuva e Cinza Crômio)"),
    ("Painel", "Painel ripado de PVC amadeirado (ACÁ), como na imagem · fixação e ponto do letreiro iluminado"),
    ("Copa", "Tijolinho Brick Invecchiatto Bianco Polare 7×23 Lepri (ACÁ), como na imagem; compatibilizar o Briefing (canjiquinha)"),
    ("Bancada", "Granito São Gabriel polido com estrutura em metal preto RAL 9005 (ACÁ)"),
    ("Marcenaria", "MDF Berneck Solanum e MDF Berneck Cinza Cristal (ACÁ: Savana + Cinza Puro); aprovar amostra antes da compra"),
    ("Piso", "Porcelanato 90×90 natural/acetinado (Briefing e ACÁ) · rodapé branco de poliestireno"),
    ("Forro", "Forro modular (Briefing e imagem) · cortineiro nas portas e janelas"),
    ("Iluminação", "Spots de embutir + pendentes sobre a mesa e a bancada (Briefing) · luz neutra 4000 K"),
    ("Mini market", "Marcenaria com fita de LED e logo iluminado; ponto para máquina de snacks · portas do rooftop na cor da parede (Indicação 03)")], PILLS30)

# ================= 20 · FICHA ROOFTOP EXTERNO =================
S = B.dup(32)
S.T(1007, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP EXTERNO")
S.T(1008, "Ficha de premissas: piscina e deck")
S.P(1012, RF(7), (0.5, 0.5)); label_w(S, 1013, "CONCEITO · PISCINA")
S.P(1014, RF(5), (0.5, 0.5)); label_w(S, 1015, "DECK + CASINHA")
S.T(1016, "Piscina de borda infinita com revestimento esverdeado, deck amadeirado, banco com floreira e logo-casinha. Indicação 01: deck.")
rows(S, ROWS32, [
    ("Piso externo", "Porcelanato que imita madeira 120×20, rodapé embutido (Briefing e Indicação 01) · peça antiderrapante"),
    ("Piscina", "Revestimento esverdeado 20×20 (Briefing) · borda infinita (imagem e memorial) · projeto próprio detalha no executivo"),
    ("Borda", "Porcelanato que imita madeira 120×20, contínuo com o deck, como na imagem; compatibilizar memorial (atérmica/São Tomé)"),
    ("Guarda-corpo", "Vidro incolor com perfil \"U\" inox (Briefing e imagem)"),
    ("Banco", "Floreira em alvenaria + banco em concreto resinado (Briefing) · tomadas junto ao banco"),
    ("Identidade", "Logo-casinha preto na parede do deck, como na imagem · fixação detalhada por Projetos"),
    ("Pergolado", "Vigas em metal preto RAL 9005, como na imagem (estilo Industrial); compatibilizar memorial (corten/madeira)"),
    ("Ducha", "Ducha + ralo e torneira próxima (Briefing) · incluir no hidrossanitário"),
    ("Aquecimento", "Solar para a piscina (Briefing) · projeto específico no executivo")], PILLS32)

# ================= 21 · PONTOS DE ATENÇÃO INTERIORES =================
S = B.dup(24)
S.T(767, "02 · INTERIORES DAS ÁREAS COMUNS · PONTOS DE ATENÇÃO")
S.T(768, "Áreas comuns: pontos de atenção na obra")
S.P(772, REC(3), (0.3, 0.5))
label_w(S, 773, "CONCEITO · JARDIM VERTICAL")
checklist(S, [
    ("Executivo de interiores", "Seguir o conceito Seazone (imagens) e os materiais do projeto básico ACÁ até a entrega do executivo ACÁ; trocas só com validação registrada."),
    ("Elétrico", "Letreiros, logo iluminado, frisos de LED do painel, fita de LED e pontos do mini market entram no elétrico antes do fechamento."),
    ("Cores de parede", "Cinza Urbano + Calopsita Suvinil (Briefing) e cimento queimado Av. Expressa: compatibilizar o ACÁ antes da compra."),
    ("Copa do rooftop", "Tijolinho Brick Bianco Polare Lepri como na imagem: compatibilizar o Briefing (canjiquinha) antes da compra."),
    ("Jardim vertical", "Preservado (ACÁ e Briefing) com o visual das imagens; detalhar fixação, impermeabilização e iluminação."),
    ("Piscina e deck", "Esverdeado 20×20 e deck em porcelanato madeira 120×20 (Briefing); aquecimento solar e ducha no executivo.")])

# ================= 22 · DIVISOR 03 =================
S = B.dup(33)
S.P(1048, RF(1), (0.5, 0.5))
S.T(1051, "Materiais de interiores das áreas comuns")
S.geom(1052, y=S.pos(1052)[1] + 0.75)
S.T(1052, "Mapa de materialidade T13–T24: amostras tiradas dos renders Seazone, do projeto básico ACÁ e padrões do catálogo Industrial para compra e aprovação")

# ================= 23-24 · MATERIAIS =================
NOTE_M = "Amostras dos renders de interiores Seazone e do projeto básico ACÁ (moodboard recepção e rooftop). Referência visual, não substitui amostra física."
for title, note, cards in [
    ("Materiais · paredes (1/3)", NOTE_M, [
        ("T13", "Cimento queimado Av. Expressa", "Paredes da recepção.",
         "Pintura cimento queimado Av. Expressa Suvinil (ACÁ) · Briefing: cimento queimado cinza.",
         "Aplicação por equipe única; aprovar painel-teste no local.", O("cimento")),
        ("T14", "Jardim vertical", "Parede da recepção, junto ao banco de madeira.",
         "Placa de jardim vertical preservado (ACÁ e Briefing), com o visual das imagens.",
         "Detalhar fixação, impermeabilização da parede e iluminação.", REC(3), (0.22, 0.5), 0.2),
        ("T15", "Tijolinho Brick Bianco Polare", "Parede da copa do rooftop.",
         "Brick Invecchiatto Bianco Polare 7×23 Lepri (ACÁ), como na imagem.",
         "Compatibilizar o Briefing (pedra canjiquinha); aprovar amostra antes da compra.", RF(1), (0.85, 0.35), 0.18),
        ("T16", "Painel ripado de PVC", "Parede do rooftop interno, com letreiro.",
         "Painel ripado de PVC amadeirado (ACÁ), como na imagem.",
         "Prever fixação e ponto do letreiro iluminado.", RF(1), (0.05, 0.66), 0.08)]),
    ("Materiais · pintura, pedra, piscina e deck (2/3)", NOTE_M, [
        ("T17", "Tintas Cinza Urbano e Calopsita", "Paredes da recepção e do rooftop.",
         "Pintura acrílica Suvinil Cinza Urbano + Calopsita (Briefing) · ACÁ: Nuvem de Chuva e Cinza Crômio.",
         "Compatibilizar o projeto ACÁ; aprovar amostra das cores antes da compra.", RF(2), (0.86, 0.62), 0.22),
        ("T18", "Granito São Gabriel polido", "Bancada da copa do rooftop.",
         "Granito São Gabriel polido (ACÁ e moodboard) com estrutura em metal preto RAL 9005.",
         "Deixar pontos hidráulicos e elétricos conforme as vistas do Briefing.", O("saogabriel")),
        ("T19", "Piscina esverdeada", "Piscina do rooftop, borda infinita.",
         "Revestimento esverdeado 20×20 (Briefing), como na imagem.",
         "Projeto próprio da piscina detalha revestimento e borda; aprovar amostra contra a imagem.", RF(7), (0.6, 0.72), 0.18),
        ("T20", "Deck porcelanato madeira", "Piso externo do rooftop e borda da piscina.",
         "Porcelanato que imita madeira 120×20 (Briefing e Indicação 01).",
         "Peça antiderrapante; caimento e paginação com a imagem.", RF(5), (0.45, 0.88), 0.14)]),
]:
    S = B.dup(34)
    S.T(1060, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
    S.T(1061, title); S.T(1065, note)
    for ids, c in zip(MC, cards):
        card(S, ids, *c)

# ================= 25 · MDFs E METAIS =================
S = B.dup(37)
S.T(1141, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
S.T(1142, "Materiais · MDFs e metais (3/3)")
S.T(1146, "Padrões Berneck e cores de metal do catálogo Industrial Seazone (rev. 15/10/2026). Referência visual, não substitui amostra física.")
for ids, c in zip(MD, [
        ("T21", "MDF Solanum", "Painel da recepção, marcenaria do rooftop e mini market, cabeceiras e aéreos.",
         "MDF Berneck Solanum, amadeirado, com perfis em metal preto RAL 9005.",
         "Compatibilizar o projeto básico ACÁ e o moodboard (Savana Guararapes); aprovar amostra física antes da compra.", swc("amostra_solanum.png")),
        ("T22", "MDF Cinza Cristal", "Marcenaria da copa do rooftop e demais marcenarias cinza.",
         "MDF Berneck Cinza Cristal (Tx), liso, em composição com o Solanum.",
         "Compatibilizar o projeto básico ACÁ e o moodboard (Cinza Puro Arauco); aprovar amostra física antes da compra.", swc("amostra_cinza_cristal.png")),
        ("T23", "Metal preto RAL 9005", "Perfis do painel, estrutura da bancada, pergolado e estruturas de mobiliário.",
         "Metal pintado preto RAL 9005 (catálogo Industrial).",
         "Mesmo acabamento em todas as peças metálicas (ACÁ: metalon preto fosco); aprovar amostra de cor.", swc("amostra_metal_preto_ral9005.png")),
        ("T24", "Metal grafite RAL 7024", "Segunda cor de metal do catálogo Industrial · referência de cor das aletas da fachada (T05).",
         "Metal pintado grafite RAL 7024 (catálogo Industrial).",
         "Usar só onde o projeto pedir segunda cor de metal.", swc("amostra_metal_grafite_ral7024.png"))]):
    card(S, ids, *c)

# ================= 26 · ENCERRAMENTO =================
S = B.dup(38)
S.P(1168, EX["E11N"], (0.5, 0.5))
S.T(1170, "MARISTA 144 SPOT")
S.T(1171, "Book de Premissas de Fachada e Interiores das Áreas Comuns")
S.geom(1172, y=S.pos(1172)[1] + 0.55)

finish(B, OUT, "MARISTA 144 SPOT")
