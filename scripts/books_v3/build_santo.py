"""Santo Antônio Spot no layout v3 (template Campeche). Fonte = book legado da Rachel (pptx) + regras v3 (15/10/2026).
python build_santo.py DADOS saida.pptx   (DADOS = trabalho/v3_dados/santo)
Estilo Industrial. Trocas aprovadas (Raquel 15/10): Jequitibá -> Carvalho Treviso · Cinza Puro -> Chumbo ·
metal preto fosco -> metal preto RAL 9005.
"""
import glob
import os
import sys

sys.path.insert(0, os.path.expanduser("~/Claude/.claude/skills/book-premissas/scripts/v3"))
from bookv3 import *  # noqa: F401,F403
from pptx.util import Pt  # noqa: E402

D, OUT = sys.argv[1], sys.argv[2]
OLD = os.path.join(REPO, "books", "legado", "SANTO_ANTONIO_SPOT_Book_de_Premissas_legado.pptx")
src = lambda n: os.path.join(D, "src", n + ".jpg")
SWC = os.path.join(os.path.dirname(os.path.abspath(D)), "sw_common")
swc = lambda n: os.path.join(SWC, n)
O = lambda k: glob.glob(os.path.join(D, "old", k + "_*"))[0]
EI = lambda n: src(f"E{n:02d}")
REC = lambda n: src(f"REC{n:02d}")
RF = lambda n: src(f"ROOF{n:02d}")
CORR = lambda n: src(f"CORR{n:02d}")
# amostras do legado (pic id = id do rotulo - 2)
PILAR, EMPENA, RIPADO, FORRO, ESQ = O("s05_12"), O("s05_19"), O("s05_26"), O("s05_33"), O("s05_40")
TORRE, MURO, PISCINA = O("s06_26"), O("s06_40"), O("s06_47")
LAJOTA, SAVANE, VELUDO = O("s17_19"), O("s17_33"), O("s17_40")
TREVISO, CHUMBO = swc("amostra_carvalho_treviso.png"), swc("amostra_chumbo.png")
PRETO, GRAFITE = swc("amostra_metal_preto_ral9005.png"), swc("amostra_metal_grafite_ral7024.png")

B = Book(TPL, os.path.join(D, "cache"))

# 1 · CAPA
S = B.dup(1)
S.P(16, EI(3), (0.5, 0.5))
S.P(17, EI(9), (0.5, 0.55))
S.P(18, REC(2), (0.4, 0.5))
S.P(19, RF(8), (0.5, 0.5))
S.T(21, "BOOK DE PREMISSAS DE FACHADA E INTERIORES DAS ÁREAS COMUNS")
S.sh(21).text_frame.paragraphs[0].runs[0].font.size = Pt(8.5)
S.T(22, "Santo Antônio Spot")
S.sh(22).text_frame.paragraphs[0].runs[0].font.size = Pt(30)
S.T(24, "Lançamento  ×  Projeto  ·  Fachada  ·  Interiores das áreas comuns  ·  Estilo Industrial")
S.T(25, "Rodovia Gilson da Costa Xavier · Santo Antônio de Lisboa · Florianópolis/SC", "DOCUMENTO TÉCNICO INTERNO")

# 2 · CONTEXTUALIZAÇÃO
S = B.dup(2)
S.T(36, "Consolidar as premissas de fachada e de interiores das áreas comuns do Santo Antônio Spot a partir da leitura cruzada "
        "entre as imagens de lançamento (R03), o anteprojeto da C2 Arquitetura, o Briefing de premissas e o projeto de interiores "
        "da Seazone: o que a imagem vendeu, o que o projeto traz, onde divergem e como a obra deve executar.")
S.P(37, EI(1), (0.42, 0.55), 0.6)
S.P(40, RIPADO)
S.P(43, EI(9), (0.45, 0.42), 0.3)
S.T(45, "CONFIRA", "o que a obra precisa observar e seguir")
for sid, q in zip([49, 51, 53, 55, 57, 59, 61], [
        "O que as imagens de lançamento mostram na fachada?",
        "O que o projeto especifica e como a obra completa?",
        "Qual é o revestimento das paredes e dos pilares?",
        "Como são o ripado, o forro de madeira e o pergolado?",
        "Qual é o conceito de interiores das áreas comuns que deve ser seguido?",
        "O que não pode ser alterado sem validação?",
        "O que a obra precisa observar na execução?"]):
    S.T(sid, q)
S.T(62, "Regra do material: nenhuma informação foi inventada. Estilo de interiores: Industrial. Prevalecem a imagem de lançamento, "
        "o conceito Seazone e o Briefing; o projeto compatibiliza. Fontes (1.27 - [3656] Santo Antônio Spot / 02 - Projetos): "
        "08 - Projeto Arquitetônico (imagens 3D R03, anteprojeto C2 rev7, Projeto Legal, interiores), 09 - Memorial Descritivo e "
        "00 - Premissas de Projeto (Briefing e Indicações). Interiores das unidades, índices urbanísticos, áreas e custos não fazem "
        "parte deste Book.")

# 3 · COMO USAR
S = B.dup(3)
S.P(75, RIPADO)
S.T(77, "T03")
S.T(78, "Código", "Código da textura (T01 a T24) ou da indicação (IND 01, IND 08) para achar o item nas premissas e no projeto.")
S.P(87, EI(1), (0.45, 0.5), None, 0.8)
S.P(90, REC(2), (0.4, 0.5))
S.P(93, RF(8), (0.5, 0.5))
S.T(89, "Fachada", "imagens de lançamento · paleta · leitura geral · texturas T01–T12 · vistas frente, conjunto e rooftop · pontos de atenção")
S.T(92, "Interiores das áreas comuns", "conceito × decisão · recepção e corredores · rooftop interno e externo · pontos de atenção")
S.T(95, "Materiais de interiores das áreas comuns", "texturas T13–T24 para compra e aprovação de amostras, com MDFs e metais do estilo Industrial")
S.T(79, "STATUS DOS ITENS")
S.T(81, "premissa validada pela Seazone: seguir na compra e na execução")
S.delete(82, 83, 84, 85)

# 4 · DIVISOR 01
S = B.dup(4)
S.P(101, EI(6), (0.5, 0.5))
S.T(105, "Imagens de lançamento  ·  paleta  ·  leitura geral  ·  texturas  ·  vistas com marcação e especificação  ·  pontos de atenção")

# 5 · MOODBOARD FACHADA
S = B.dup(5)
S.P(118, EI(1), (0.48, 0.5)); label_w(S, 119, "01 · FRENTE")
S.P(120, EI(9), (0.5, 0.55)); label_w(S, 121, "09 · CONJUNTO")
S.P(122, EI(3), (0.5, 0.5)); label_w(S, 123, "03")
S.P(124, EI(7), (0.5, 0.5)); label_w(S, 125, "07")
S.P(126, EI(5), (0.35, 0.5)); label_w(S, 127, "05 · ROOFTOP")
S.P(128, EI(8), (0.65, 0.6)); label_w(S, 129, "08")
for sid, (a, b) in zip([132, 134, 136, 138, 140], [
        ("Pilares claros", "placas claras com aspecto de pedra emoldurando a frente"),
        ("Madeira na frente", "ripado vertical em frente às janelas e forro amadeirado com spots sob as lajes"),
        ("Vidro e preto", "guarda-corpos de vidro e esquadrias pretas de caixilho quadriculado"),
        ("Conjunto em patamares", "volumes claros escalonados no terreno e torres de circulação grafite com logo Seazone"),
        ("Rooftop", "pergolado de madeira com vegetação, floreira contínua e piscina esverdeada voltada para a baía")]):
    S.T(sid, a, b)

# 6 · PALETA
S = B.dup(6)
for pic, lab, txt, code, a, b, img, foc, hf in [
        (151, 152, 153, "T01", "Placas claras (pedra)", "pilares da frente", PILAR, (0.5, 0.5), None),
        (154, 155, 156, "T02", "Concreto cinza texturizado", "empenas e paredes laterais", EMPENA, (0.5, 0.8), 0.5),
        (157, 158, 159, "VOLUMES", "Papel de Presente B208", "volumes claros do conjunto", EI(9), (0.79, 0.74), 0.07),
        (160, 161, 162, "T03", "Brise alumínio amadeirado", "frente das janelas", RIPADO, (0.5, 0.5), None),
        (163, 164, 165, "T04", "Forro amadeirado", "face inferior das lajes", FORRO, (0.5, 0.5), None),
        (166, 167, 168, "T09", "Torre grafite", "circulação vertical com logo", TORRE, (0.2, 0.5), 0.6)]:
    S.P(pic, img, foc, hf); S.T(lab, code); S.T(txt, a, b)
S.T(170, "COMO LER", "Fachada = pilares em placas claras (T01) + empenas em concreto cinza (T02) + brise em alumínio amadeirado (T03) "
                     "+ forro amadeirado com spots (T04) + esquadria preta + guarda-corpo de vidro. No conjunto, volumes claros em Papel "
                     "de Presente Suvinil (B208) e torres grafite (T09). Detalhe de cada material nas páginas de texturas T01–T12.")

# 7 · LEITURA GERAL
S = B.dup(7)
win = S.P(181, EI(1))
label_w(S, 182, "LANÇAMENTO · 01 (FRENTE)")
place_markers(S, 181, win, list(range(183, 193)), old_markers(OLD, 4))
S.T(193, "A frente vendida tem três pavimentos com pilares claros, ripado e forro de madeira sob as lajes, guarda-corpos de vidro e "
         "esquadrias pretas. O térreo é envidraçado, com o letreiro SANTO ANTÔNIO SPOT sobre a marquise, e o rooftop tem floreira "
         "contínua. O anteprojeto não traz legenda de material: vale a imagem e o executivo compatibiliza.")
for sid, (a, b) in zip(range(196, 215, 2), [
        ("Floreira do rooftop", "vegetação contínua na borda"),
        ("Guarda-corpo do rooftop", "vidro com perfil \"U\" inox (Briefing)"),
        ("Forro de madeira", "amadeirado sob as lajes, com spots"),
        ("Ripado vertical", "brise em alumínio amadeirado"),
        ("Pilares em placas claras", "aspecto de pedra, como na imagem"),
        ("Guarda-corpo das sacadas", "vidro + alumínio anodizado (Briefing)"),
        ("Esquadrias pretas", "linha 32 (gold), caixilho quadriculado"),
        ("Letreiro", "SANTO ANTÔNIO SPOT em relevo"),
        ("Paredes em concreto", "textura cinza nas empenas"),
        ("Muro", "textura rústica cinza + logo e nome")]):
    S.T(sid, a, b)

# 8-10 · TEXTURAS DA FACHADA
NOTE = ("Amostras extraídas das imagens de lançamento (R03): referência visual, não substitui amostra física. Especificação: "
        "Briefing, Memorial Descritivo e anteprojeto C2 rev7; onde o projeto não traz legenda de material, vale a imagem.")
for title, cards in [
    ("Texturas da fachada · pilares, paredes e madeira (1/3)", [
        ("T01", "Pilares em placas claras", "Pilares e montantes da fachada frontal.",
         "Placas claras com veios, aspecto de pedra, como na imagem de lançamento (memorial: revestimento conforme projeto).",
         "Projetos detalha material e paginação no executivo; aprovar amostra física contra a imagem antes da compra.", PILAR),
        ("T02", "Concreto cinza texturizado", "Empenas e paredes laterais.",
         "Textura efeito concreto cinza, como na imagem (Briefing: área externa com pintura seguindo o conceito da fachada).",
         "Painel-teste contra a imagem antes da compra; mesma cor e lote em todas as faces.", EMPENA, (0.5, 0.85)),
        ("T03", "Ripado vertical", "Frente das janelas, em todos os pavimentos.",
         "Brise em alumínio com acabamento amadeirado, ripas verticais como na imagem.",
         "Anteprojeto e memorial indicam madeira: executar em alumínio, acabamento amadeirado no tom da imagem; aprovar amostra antes da compra.", RIPADO),
        ("T04", "Forro de madeira", "Face inferior das lajes em balanço.",
         "Forro de acabamento amadeirado, resistente a intempérie, com spots embutidos, como na imagem.",
         "Detalhar material no executivo; aprovar amostra contra a imagem; spots no luminotécnico.", FORRO)]),
    ("Texturas da fachada · esquadrias e rooftop (2/3)", [
        ("T05", "Esquadrias pretas", "Térreo, áreas comuns e janelas da fachada.",
         "Linha 32 (gold) com perfil preto; áreas comuns em preto fosco (Briefing) · caixilho quadriculado como na imagem.",
         "Converge com a imagem: indicar linha e cor no quadro de esquadrias antes do pedido.", ESQ),
        ("T06", "Guarda-corpos", "Sacadas e rooftop.",
         "Sacadas: vidro com alumínio anodizado de montante aparente (Briefing) · rooftop: vidro incolor com perfil \"U\" inox (Briefing).",
         "Detalhar perfil e fixação no executivo antes do pedido.", EI(3), (0.27, 0.64), 0.12),
        ("T07", "Pergolado de madeira", "Rooftop, sobre a piscina.",
         "Pilares e vigas de madeira com vegetação, como na imagem (memorial admite ripas de madeira).",
         "Não executar em aço corten; Projetos detalha estrutura, seção e tratamento da madeira no executivo.", EI(5), (0.40, 0.36), 0.16),
        ("T08", "Floreira do rooftop", "Borda do último pavimento.",
         "Floreira contínua em alvenaria com vegetação, como na imagem (Briefing: floreira em alvenaria).",
         "Impermeabilização e irrigação no executivo; paisagismo incluído no complementar.", EI(5), (0.68, 0.64), 0.16)]),
    ("Texturas da fachada · identidade, muro e piscina (3/3)", [
        ("T09", "Torre escura com logo", "Volume da circulação vertical.",
         "Torre grafite (anteprojeto: cinza escuro) com logo Seazone iluminado, como na imagem do conjunto.",
         "Projetos detalha fixação e ponto elétrico do logo antes do fechamento da fachada.", TORRE),
        ("T10", "Letreiro SANTO ANTÔNIO SPOT", "Marquise do térreo.",
         "Letras brancas em relevo, como na imagem (memorial: nome do edifício na fachada).",
         "Como na imagem; Projetos detalha fixação e ponto elétrico; luminotécnico antes do fechamento de fachada.", EI(3), (0.6, 0.645), 0.12),
        ("T11", "Muro em textura rústica", "Fechamento frontal.",
         "Muro cinza em textura rústica, como na imagem, com logo e nome Seazone (Briefing).",
         "Prever fixação e iluminação do logo; aprovar amostra da textura contra a imagem.", MURO),
        ("T12", "Piscina do rooftop", "Rooftop com vista para a baía.",
         "Revestimento esverdeado 20×20 (Briefing) · lajota verde na imagem e nas texturas do projeto de interiores.",
         "Converge com a imagem: aprovar amostra antes do pedido; aquecimento solar (Briefing).", PISCINA)])]:
    S = B.dup(8)
    S.T(221, title); S.T(225, NOTE)
    for ids, c in zip(TEX, cards):
        card(S, ids, *c)

# 11 · VISTA FRENTE (tpl 12, vertical)
S = B.dup(12)
S.T(354, "Frente e acesso")
win = S.P(358, EI(3))
label_w(S, 359, "LANÇAMENTO · 03 (FRENTE)")
place_markers(S, 358, win, list(range(360, 369)), old_markers(OLD, 7))
S.T(370, LEG)
vista_items(S, V12, [
    ("Pilares", "Placas claras com aspecto de pedra, como na imagem"),
    ("Forro", "Amadeirado sob as lajes, com spots embutidos"),
    ("Ripado", "Brise em alumínio amadeirado, como na imagem"),
    ("Empena", "Textura efeito concreto cinza, como na imagem"),
    ("Guarda-corpo", "Vidro + alumínio anodizado de montante aparente (Briefing)"),
    ("Letreiro", "SANTO ANTÔNIO SPOT em letras brancas em relevo"),
    ("Esquadrias", "Linha 32 (gold), perfil preto (Briefing)"),
    ("Muro", "Textura rústica cinza + logo e nome Seazone (Briefing)")])
S.P(398, PILAR); S.P(399, RIPADO); S.P(400, FORRO)
S.T(401, "T01 · T03 · T04")
S.T(403, "SEM QUADRO DE MATERIAIS", "O anteprojeto rev7 só indica \"madeira\" nos brises e o Projeto Legal não tem elevações: vale a "
         "imagem de venda. O executivo de arquitetura traz o quadro de materiais da fachada seguindo este book.")

# 12 · VISTA CONJUNTO (tpl 11)
S = B.dup(11)
S.T(302, "Vista do conjunto")
win = S.P(306, EI(9))
label_w(S, 307, "LANÇAMENTO · 09 (CONJUNTO)")
place_markers(S, 306, win, list(range(308, 317)), old_markers(OLD, 8))
S.T(318, "LEITURA DO CONJUNTO", "O conjunto se escalona no terreno em volumes claros, com duas torres de circulação escuras e muro de "
         "contenção de concreto no talude. O anteprojeto (elevação noroeste) mostra a mesma composição, com as torres em cinza "
         "escuro e os volumes claros.")
S.T(320, LEG)
vista_items(S, V11, [
    ("Torre", "Grafite com logo Seazone iluminado"),
    ("Volumes", "Papel de Presente Suvinil (B208), em patamares"),
    ("Esquadrias", "Linha 32 (gold), perfil preto"),
    ("Contenção", "Muro de concreto no talude, como na imagem"),
    ("Pergolado", "Madeira com vegetação no rooftop"),
    ("Floreira", "Alvenaria com vegetação contínua no rooftop"),
    ("Letreiro", "SANTO ANTÔNIO SPOT na frente da rodovia"),
    ("Talude", "Gramado, como na imagem (paisagismo no complementar)")])

# 13 · VISTA ROOFTOP (tpl 13) — marcadores posicionados sobre a imagem 05
S = B.dup(13)
S.T(410, "Rooftop: pergolado, floreira e piscina")
win = S.P(414, EI(5))
label_w(S, 415, "LANÇAMENTO · 05 (ROOFTOP)")
place_markers(S, 414, win, list(range(416, 426)), [
    (0.20, 0.30), (0.25, 0.13), (0.47, 0.57), (0.33, 0.62), (0.68, 0.66), (0.38, 0.84), (0.63, 0.88), (0.555, 0.72)])
S.T(427, "ATENÇÃO · ROOFTOP", "Pergolado de madeira com vegetação, piscina esverdeada e floreira contínua são a imagem do rooftop no "
         "lançamento: executar como na imagem. Projetos detalha a estrutura do pergolado, a impermeabilização e a irrigação no executivo.")
S.T(429, LEG)
vista_items(S, V13, [
    ("Pergolado", "Pilares e vigas de madeira, como na imagem"),
    ("Vegetação", "Trepadeira sobre o pergolado, como na imagem"),
    ("Guarda-corpo", "Vidro incolor com perfil \"U\" inox (Briefing)"),
    ("Piscina", "Revestimento esverdeado 20×20 · aquecimento solar"),
    ("Floreira", "Alvenaria com vegetação, impermeabilizada"),
    ("Volumes claros", "Papel de Presente Suvinil (B208)"),
    ("Esquadrias", "Linha 32 (gold), perfil preto"),
    ("Pilares e muretas", "Claros, como na imagem")])

# 14 · PONTOS DE ATENÇÃO FACHADA
S = B.dup(24)
S.T(767, "01 · FACHADA · PONTOS DE ATENÇÃO")
S.T(768, "Fachada: pontos de atenção na obra")
S.P(772, EI(1), (0.48, 0.5))
label_w(S, 773, "LANÇAMENTO · 01")
checklist(S, [
    ("Revestimentos da frente", "Pilares em placas claras e empenas em concreto cinza, como na imagem: Projetos detalha no executivo; amostra aprovada antes da compra."),
    ("Brises", "Alumínio com acabamento amadeirado, ripas como na imagem: compatibilizar anteprojeto e memorial (madeira) antes de fabricar."),
    ("Guarda-corpos", "Sacadas em vidro com alumínio anodizado de montante aparente; rooftop em vidro com perfil \"U\" inox (Briefing)."),
    ("Pergolado e forro", "Madeira, como na imagem (não executar em aço corten); forro amadeirado resistente a intempérie no executivo."),
    ("Letreiro e logo", "Como na imagem; Projetos detalha fixação e ponto elétrico; luminotécnico entregue antes do fechamento de fachada."),
    ("Muro e paisagismo", "Muro rústico com logo e nome Seazone; floreiras, talude e contenção incluídos no complementar.")])

# 15 · DIVISOR 02
S = B.dup(14)
S.P(465, REC(1), (0.5, 0.5))
S.T(468, "Interiores das áreas comuns")
S.T(469, "Conceito Seazone × decisão  ·  recepção e corredores  ·  rooftop interno e mini market  ·  rooftop externo e piscina  ·  pontos de atenção")

# 16 · CONCEITO × DECISÃO
S = B.dup(15)
S.T(477, "02 · INTERIORES DAS ÁREAS COMUNS")
S.T(478, "Conceito Seazone × decisão registrada")
concept_cards(S, [
    ("Recepção", REC(1), (0.45, 0.5),
     "Jardim vertical com logo, painel ripado claro, sofá em \"L\", balanço de fibra, espelho orgânico, frase na parede",
     "Textura velvet + painel PVC ripado + jardim vertical (Briefing)"),
    ("Corredores", CORR(1), (0.5, 0.5),
     "Paredes claras, portas pretas com número e arandela em cada porta",
     "Cinza Urbano + Calopsita (Briefing), em tons claros como na imagem"),
    ("Rooftop interno", RF(2), (0.5, 0.5),
     "Mini market, copa com bancada escura, painel ripado com logo, mesa redonda",
     "Textura velvet + painel PVC ripado (Briefing) · pontos do mini market no elétrico"),
    ("Rooftop externo", RF(8), (0.5, 0.5),
     "Pergolado de madeira com vegetação, piscina com lajota verde, deck, banco com almofadas, redes",
     "Piscina esverdeada 20×20 · porcelanato imitando madeira 120×20 só na área da piscina (Briefing)"),
    ("Projeto de interiores", RF(13), (0.5, 0.5),
     "Imagens, texturas e moodboard final Seazone (10 - Projeto de Interiores)",
     "Executivo de interiores: seguir o conceito Seazone (imagens) até a entrega do executivo")])
S.T(507, "Imagens de interiores Seazone (10 - Projeto de Interiores / 02 - Imagens) × Briefing, Memorial, moodboard e texturas do projeto de interiores.")
S.T(509, "ATENÇÃO · REGRA PARA A OBRA", "Vale o conceito das imagens de venda, o moodboard final e as texturas da pasta \"01 - Arquivos de "
         "projeto / Texturas\". Logos, frases, jardim vertical, balanço e pontos do mini market entram no elétrico, no hidráulico e no "
         "estrutural antes do fechamento.")

# 17 · FICHA RECEPÇÃO
S = B.dup(27)
S.T(848, "02 · INTERIORES DAS ÁREAS COMUNS · RECEPÇÃO E CORREDORES")
S.T(849, "Ficha: recepção, hall e corredores")
S.P(853, REC(2), (0.45, 0.5)); label_w(S, 854, "CONCEITO SEAZONE · RECEPÇÃO")
S.P(855, CORR(1), (0.5, 0.5)); label_w(S, 856, "CONCEITO · CORREDOR")
S.T(857, "Jardim vertical com logo, painel ripado claro, sofá em \"L\" e balanço de fibra com a frase \"Aqui, até o pôr do sol é mais bonito\".")
rows(S, ROWS27, [
    ("Parede-destaque", "Jardim vertical com logo Santo Antônio SPOT iluminado (Briefing e imagem) · ponto elétrico e irrigação"),
    ("Paredes", "Pintura textura velvet (Briefing) · textura velvet Coral (moodboard de interiores)"),
    ("Painel", "Painel PVC ripado (Briefing), ripado claro como na imagem"),
    ("Forro", "Gesso com pé-direito livre de 2,60 m (Briefing)"),
    ("Piso", "Porcelanato 90×90 natural/acetinado (Briefing) · rodapé de poliestireno preto"),
    ("Balanço", "Poltrona balanço de fibra presa na laje, com reforço previsto no estrutural (Briefing)"),
    ("Portas", "Portas e rodapés das áreas comuns em preto, bandeira até o forro (Briefing)"),
    ("Mobiliário", "Sofá em \"L\", mesas de apoio, espelho orgânico e aparador (imagem)"),
    ("Corredores", "Cinza Urbano + Calopsita (Briefing), tons claros como na imagem · arandelas nas portas")], PILLS27)

# 18 · MOODBOARD ROOFTOP (tpl 31)
S = B.dup(31)
S.T(977, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP · MOODBOARD")
S.T(978, "Rooftop: mini market, copa, terraço e piscina")
for pic, img, foc in zip(range(982, 993),
                         [RF(2), RF(5), RF(3), RF(4), RF(8), EI(7), RF(10), RF(13), RF(11), RF(15), RF(9)],
                         [(.5, .5), (.5, .5), (.5, .5), (.5, .5), (.5, .5), (.5, .6), (.5, .5), (.5, .5), (.5, .5), (.5, .5), (.5, .5)]):
    S.P(pic, img, foc)
S.T(993, "CONCEITO", "Mini market com estante de produtos, copa com bancada escura, painel ripado com o logo Santo Antônio SPOT, "
         "mesa redonda e poltronas claras.",
         "TERRAÇO", "Pergolado de madeira com vegetação sobre a piscina com lajota verde, banco com almofadas, redes, deck "
         "amadeirado e guarda-corpo de vidro.",
         "ESTILO", "Industrial: MDF Carvalho Treviso e Chumbo, metal preto RAL 9005.")
for pic, lab, name, img, foc, hf in [(994, 995, "Lajota", LAJOTA, (.5, .5), None),
                                    (996, 997, "Velvet", REC(3), (0.86, 0.35), 0.15),
                                    (998, 999, "Ripado", REC(2), (0.50, 0.35), 0.15),
                                    (1000, 1001, "Treviso", TREVISO, (.5, .5), None)]:
    S.P(pic, img, foc, hf); S.T(lab, name)

# 19 · FICHA ROOFTOP INTERNO
S = B.dup(30)
S.T(938, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP INTERNO")
S.T(939, "Ficha de premissas: copa e mini market")
S.P(943, RF(2), (0.5, 0.5)); label_w(S, 944, "CONCEITO · COPA E MINI MARKET")
S.P(945, RF(5), (0.5, 0.5)); label_w(S, 946, "PAINEL RIPADO COM LOGO")
S.T(947, "Mini market com estante de produtos e copa com bancada escura na área coberta; painel ripado com o logo do Spot.")
rows(S, ROWS30, [
    ("Paredes", "Pintura textura velvet + painel PVC ripado (Briefing)"),
    ("Painel", "Painel ripado com logo Santo Antônio SPOT, como na imagem · ponto elétrico do logo"),
    ("Forro", "Forro modular 2,60 m (Briefing) · cortineiro de 20 cm nas portas e janelas"),
    ("Piso", "Porcelanato 90×90 natural/acetinado (Briefing) · Urban Savane (texturas)"),
    ("Mini market", "Estante com produtos (imagem) · pontos elétricos da estante no projeto elétrico"),
    ("Copa", "Bancada escura, como na imagem · pontos hidráulicos e elétricos no complementar"),
    ("Portas", "Portas e rodapés das áreas comuns em preto (Briefing)"),
    ("Esquadrias", "Linha 32 (gold), acabamento preto fosco nas áreas comuns (Briefing)"),
    ("Mobiliário", "Mesa redonda, cadeiras e poltronas claras, como na imagem")], PILLS30)

# 20 · FICHA ROOFTOP EXTERNO
S = B.dup(32)
S.T(1007, "02 · INTERIORES DAS ÁREAS COMUNS · ROOFTOP EXTERNO")
S.T(1008, "Ficha de premissas: rooftop externo e piscina")
S.P(1012, RF(10), (0.5, 0.5)); label_w(S, 1013, "CONCEITO · PISCINA")
S.P(1014, RF(8), (0.5, 0.5)); label_w(S, 1015, "PERGOLADO + REDES")
S.T(1016, "Pergolado de madeira com vegetação sobre a piscina, banco com almofadas, redes e vista para a baía.")
rows(S, ROWS32, [
    ("Piscina", "Revestimento esverdeado 20×20 (Briefing) · lajota verde (imagem e texturas)"),
    ("Aquecimento", "Solar para a piscina (Briefing) · incluir no hidrossanitário"),
    ("Piso externo", "Porcelanato imitando madeira 120×20 só na área da piscina · 90×90 no restante (Briefing)"),
    ("Pergolado", "Madeira com vegetação, como na imagem (memorial admite ripas de madeira)"),
    ("Guarda-corpo", "Vidro incolor com perfil \"U\" inox (Briefing)"),
    ("Banco", "Banco em concreto resinado junto à piscina, sob o pergolado (Briefing · Indicação 08)"),
    ("Floreiras", "Floreira em alvenaria (Briefing · Indicações 01 e 08) · impermeabilização e irrigação"),
    ("Redes", "Redes sob o pergolado, como na imagem · fixação detalhada no executivo"),
    ("Mobiliário", "Banco com almofadas brancas, puffs e mesas de apoio, como na imagem")], PILLS32)

# 21 · PONTOS DE ATENÇÃO INTERIORES
S = B.dup(24)
S.T(767, "02 · INTERIORES DAS ÁREAS COMUNS · PONTOS DE ATENÇÃO")
S.T(768, "Áreas comuns: pontos de atenção na obra")
S.P(772, REC(4), (0.25, 0.5))
label_w(S, 773, "CONCEITO · BALANÇO E FRASE")
checklist(S, [
    ("Executivo de interiores", "Seguir o conceito Seazone (imagens, moodboard final e texturas) até a entrega do executivo."),
    ("Logos, frases e jardim vertical", "Prever fixação, irrigação, impermeabilização e pontos elétricos antes do fechamento; luminotécnico por Projetos."),
    ("Balanço da recepção", "Poltrona balanço presa na laje (Briefing): prever o reforço no projeto estrutural."),
    ("Mini market e copa", "Pontos da estante e da copa incluídos no elétrico e no hidrossanitário."),
    ("Piscina", "Esverdeado 20×20 (Briefing), aquecimento solar; porcelanato imitando madeira 120×20 só na área da piscina."),
    ("Corredores e portas", "Cinza Urbano + Calopsita em tons claros (painel-teste antes da compra); portas e rodapés pretos.")])

# 22 · DIVISOR 03
S = B.dup(33)
S.P(1048, RF(13), (0.5, 0.5))
S.T(1051, "Materiais de interiores das áreas comuns")
S.geom(1052, y=S.pos(1052)[1] + 0.75)
S.T(1052, "Mapa de materialidade T13–T24: amostras das texturas e renders Seazone e padrões do catálogo Industrial para compra e aprovação")

# 23-24 · MATERIAIS
NOTE_M = ("Texturas da pasta \"10 - Projeto de Interiores / 01 - Arquivos de projeto / Texturas\", moodboard final e renders Seazone. "
          "Referência visual, não substitui amostra física.")
for title, cards in [
    ("Materiais · paredes e pisos (1/3)", [
        ("T13", "Textura velvet", "Paredes da recepção e do rooftop.",
         "Pintura textura velvet (Briefing) · textura velvet Coral (moodboard de interiores).",
         "Aprovar painel-teste no local antes da compra.", REC(3), (0.86, 0.35), 0.15),
        ("T14", "Painel PVC ripado", "Recepção e rooftop (parede do logo).",
         "Painel PVC ripado (Briefing), ripado claro como na imagem.",
         "Prever ponto do logo iluminado antes do painel; aprovar amostra de cor.", REC(2), (0.50, 0.35), 0.15),
        ("T15", "Jardim vertical", "Parede-destaque da recepção, com logo.",
         "Jardim vertical com logo Santo Antônio SPOT (Briefing e imagem).",
         "Detalhar fixação, irrigação, impermeabilização e iluminação do logo.", REC(2), (0.08, 0.45), 0.2),
        ("T16", "Porcelanato Urban Savane", "Pisos internos das áreas comuns.",
         "Urban Savane 90×90 (texturas) · Briefing: porcelanato 90×90 natural/acetinado.",
         "Converge com o Briefing: aprovar amostra antes da compra.", SAVANE)]),
    ("Materiais · piscina e mobiliário (2/3)", [
        ("T17", "Lajota da piscina", "Revestimento da piscina do rooftop.",
         "Textura \"Lajota piscina\" (projeto de interiores) · Briefing: esverdeado 20×20.",
         "Converge com o Briefing e a imagem: aprovar amostra antes do pedido.", LAJOTA),
        ("T18", "Porcelanato imitando madeira", "Área da piscina (deck e borda).",
         "Porcelanato imitando madeira 120×20 só na área da piscina (Briefing), visual de deck da imagem.",
         "Peça antiderrapante; conferir paginação com a imagem e aprovar amostra.", RF(11), (0.62, 0.9), 0.14),
        ("T19", "Veludo liso cinza", "Estofados das áreas comuns.",
         "Textura \"Tecido veludo liso cinza\" do projeto de interiores.",
         "Referência de mobiliário: aprovar amostra do tecido antes da compra.", VELUDO),
        ("T20", "Balanço de fibra", "Recepção, junto à frase na parede.",
         "Poltrona balanço de fibra presa na laje (Briefing e imagem).",
         "Reforço na laje previsto no projeto estrutural antes da instalação.", REC(4), (0.21, 0.72), 0.28)])]:
    S = B.dup(34)
    S.T(1060, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
    S.T(1061, title); S.T(1065, NOTE_M)
    for ids, c in zip(MC, cards):
        card(S, ids, *c)

S = B.dup(37)
S.T(1141, "03 · INTERIORES DAS ÁREAS COMUNS · MAPA DE MATERIALIDADE")
S.T(1142, "Materiais · MDFs e metais (3/3)")
S.T(1146, "Padrões Berneck e cores de metal do catálogo Industrial Seazone (rev. 13/10/2026). Referência visual, não substitui amostra física.")
for ids, c in zip(MD, [
        ("T21", "MDF Carvalho Treviso", "Marcenaria amadeirada (moodboard de interiores Santo Antônio).",
         "MDF Berneck Carvalho Treviso (Design), amadeirado claro.",
         "Compatibilizar memorial de interiores (Jequitibá); aprovar amostra física antes da compra.", TREVISO),
        ("T22", "MDF Chumbo", "Marcenaria lisa, em composição com o Carvalho Treviso.",
         "MDF Berneck Chumbo (Micro), liso.",
         "Compatibilizar memorial de interiores (Cinza Puro); aprovar amostra física antes da compra.", CHUMBO),
        ("T23", "Metal preto RAL 9005", "Espelhos, estruturas de mobiliário e metais da marcenaria.",
         "Metal pintado preto RAL 9005 (catálogo Industrial).",
         "Compatibilizar moodboard (metal preto fosco); mesmo acabamento em todas as peças; aprovar amostra de cor.", PRETO),
        ("T24", "Metal grafite RAL 7024", "Segunda cor de metal do catálogo Industrial. Sem uso definido no projeto atual.",
         "Metal pintado grafite RAL 7024 (catálogo Industrial).",
         "Usar só se o projeto pedir segunda cor de metal.", GRAFITE)]):
    card(S, ids, *c)

# 26 · ENCERRAMENTO
S = B.dup(38)
S.P(1168, EI(1), (0.48, 0.5))
S.T(1170, "SANTO ANTÔNIO SPOT")
S.T(1171, "Book de Premissas de Fachada e Interiores das Áreas Comuns")
S.T(1172, "Lançamento × Projeto  ·  Seazone")
S.geom(1172, y=S.pos(1172)[1] + 0.55)

finish(B, OUT, "SANTO ANTÔNIO SPOT")
